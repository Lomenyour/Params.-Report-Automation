"""Снимок листа поправок: значения, оформление и собственные диаграммы.

Из исходного ZIP изолируется один лист, поэтому миллион строк основного
листа не загружается в объекты openpyxl. Формулы поправок фиксируются по
сохранённым результатам: ссылки на удалённые отчётные листы не переносятся.
"""
from __future__ import annotations

from copy import copy, deepcopy
from dataclasses import dataclass
import io
import json
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.cell import WriteOnlyCell
import pandas as pd

from gas_passport import MARKER, write_passport
from correction_profiles import MARKER as CORRECTION_MARKER, FIELDS, write_profile

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@dataclass(frozen=True)
class CorrectionSheetSnapshot:
    content: bytes


def read_correction_sheet(path):
    with ZipFile(path) as source:
        workbook = ET.fromstring(source.read("xl/workbook.xml"))
        sheets = workbook.find(f"{{{NS}}}sheets")
        selected = next((s for s in sheets if s.get("name") == "Поправки"), None)
        if selected is None:
            return None
        relation_id = selected.get(f"{{{REL}}}id")
        relations = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        relation = next(r for r in relations if r.get("Id") == relation_id)
        target = relation.get("Target")
        sheet_path = target.lstrip("/") if target.startswith("/") else str(PurePosixPath("xl") / target)
        # Совместимость с ранее созданными книгами: старый экспорт оставлял
        # висячие ссылки на стили строк/колонок. Удаляется только недоступное
        # оформление в памяти, значения и файл-источник не меняются.
        xml = ET.fromstring(source.read(sheet_path))
        styles = ET.fromstring(source.read("xl/styles.xml"))
        style_count = len(styles.find(f"{{{NS}}}cellXfs"))
        repaired = 0
        for tag, attr in (("col", "style"), ("row", "s")):
            for node in xml.iter(f"{{{NS}}}{tag}"):
                if attr in node.attrib and int(node.get(attr)) >= style_count:
                    del node.attrib[attr]
                    repaired += 1
        if repaired:
            print(f"WARNING: восстановлены {repaired} ссылок оформления строк/колонок «Поправки». "
                  "Недоступный стиль заменён оформлением по умолчанию; данные не изменены.")
        for sheet in list(sheets):
            if sheet is not selected:
                sheets.remove(sheet)
        names = workbook.find(f"{{{NS}}}definedNames")
        if names is not None:
            workbook.remove(names)
        views = workbook.find(f"{{{NS}}}bookViews")
        if views is not None:
            for view in views:
                view.set("activeTab", "0")
        output = io.BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as isolated:
            for name in source.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml") and name != sheet_path:
                    continue
                content = (ET.tostring(workbook) if name == "xl/workbook.xml" else
                           ET.tostring(xml) if name == sheet_path and repaired else source.read(name))
                isolated.writestr(name, content)
    with pd.ExcelFile(path, engine="calamine") as book:
        cached = book.parse("Поправки", header=None)
    book = load_workbook(io.BytesIO(output.getvalue()), keep_links=False)
    sheet = book["Поправки"]
    # Использованные профили заменяются из неизменяемого снимка расчёта.
    from openpyxl.styles import Font, PatternFill
    starts = []
    for marker, base, key in ((MARKER, 10, "components"), (CORRECTION_MARKER, 5, "coefficients")):
        positions = [cell.row for cell in list(sheet._cells.values()) if cell.column == 1 and cell.value == marker]
        if len(positions) > 1:
            raise ValueError("Несколько блоков одного профиля на листе «Поправки».")
        for start in positions:
            starts.append(start)
            data = json.loads(sheet.cell(start, 2).value)
            end = start + base + len(data.get(key, {})) - 1
            sheet.row_dimensions[start].hidden = False
            for row in sheet.iter_rows(min_row=start, max_row=end, max_col=3):
                for cell in row:
                    cell.value = None
                    cell.font, cell.fill = Font(), PatternFill()
            for row in range(start,end+1):
                if row in sheet.row_dimensions and not any(
                    cell.row == row and cell.value is not None for cell in sheet._cells.values()):
                    del sheet.row_dimensions[row]
    if starts and not any(cell.value is not None and cell.row >= min(starts) for cell in sheet._cells.values()):
        sheet.delete_rows(min(starts), sheet.max_row - min(starts) + 1)
    for cell in list(sheet._cells.values()):
        if cell.data_type == "f":
            value = cached.iat[cell.row - 1, cell.column - 1] if (
                cell.row <= len(cached) and cell.column <= len(cached.columns)) else None
            cell.value = None if value is None or pd.isna(value) else value
    snapshot = io.BytesIO()
    book.save(snapshot)
    book.close()
    return CorrectionSheetSnapshot(snapshot.getvalue())


def _dimension_for_target(original, target):
    dimension = copy(original)
    dimension.parent = target
    dimension._style = None
    if original.has_style:
        for attr in ("font", "fill", "border", "alignment", "protection", "number_format"):
            setattr(dimension, attr, copy(getattr(original, attr)))
    return dimension


def write_correction_sheet(book, snapshot, passport=None, correction_profile=None):
    if snapshot is None and passport is None and correction_profile is None:
        return
    target = book.create_sheet("Поправки")
    if snapshot is not None:
        source_book = load_workbook(io.BytesIO(snapshot.content), keep_links=False)
        try:
            source = source_book["Поправки"]
            for name in ("sheet_format", "sheet_properties", "sheet_state", "sheet_view",
                         "sheet_protection", "page_setup", "page_margins", "print_options"):
                # sheet_view/protection имеют особые имена в openpyxl.
                if name in ("sheet_view", "sheet_protection"):
                    continue
                if hasattr(source, name):
                    setattr(target, name, copy(getattr(source, name)))
            target.views = deepcopy(source.views)
            target.protection = copy(source.protection)
            target.freeze_panes = source.freeze_panes
            target.auto_filter = deepcopy(source.auto_filter)
            target.conditional_formatting = deepcopy(source.conditional_formatting)
            target.data_validations = deepcopy(source.data_validations)
            target.print_options = copy(source.print_options)
            if source.print_area:
                target.print_area = str(source.print_area)
            target.print_title_rows = source.print_title_rows
            target.print_title_cols = source.print_title_cols
            for key, value in source.column_dimensions.items():
                target.column_dimensions[key] = _dimension_for_target(value, target)
            for key, value in source.row_dimensions.items():
                target.row_dimensions[key] = _dimension_for_target(value, target)
            last_row = max(source.max_row, max(source.row_dimensions,default=0))
            _profile_dimensions(target, last_row, correction_profile, passport)
            for merged in source.merged_cells.ranges:
                target.merged_cells.add(str(merged))
            for row in source.iter_rows(max_row=last_row):
                cells = []
                for original in row:
                    if original.value is None and not original.has_style and not original.comment and not original.hyperlink:
                        cells.append(None)
                        continue
                    cell = WriteOnlyCell(target, original.value)
                    for attr in ("font", "fill", "border", "alignment", "protection", "number_format"):
                        setattr(cell, attr, copy(getattr(original, attr)))
                    if original.comment:
                        cell.comment = copy(original.comment)
                    if original.hyperlink:
                        cell.hyperlink = copy(original.hyperlink)
                    if original.data_type == "s":
                        cell.data_type = "s"
                    cells.append(cell)
                target.append(cells)
            for chart in source._charts:
                target.add_chart(deepcopy(chart))
            # Изображения уже загружены openpyxl из изолированной книги.
            for picture in source._images:
                target.add_image(deepcopy(picture))
        finally:
            source_book.close()
    if snapshot is None:
        _profile_dimensions(target, 0, correction_profile, passport)
    if correction_profile is not None:
        target.append([])
        target.append([])
        write_profile(target, correction_profile)
    if passport is not None:
        target.append([])
        target.append([])
        write_passport(target, passport)


def _profile_dimensions(sheet, base_row, corrections, passport):
    cursor = base_row
    if corrections is not None:
        start = cursor + 3
        sheet.row_dimensions[start].hidden = True
        for row in range(start + 4, start + 5 + len(FIELDS)):
            sheet.row_dimensions[row].height = 48
        cursor += 2 + 5 + len(FIELDS)
    if passport is not None:
        sheet.row_dimensions[cursor + 3].hidden = True
