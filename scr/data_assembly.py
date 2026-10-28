"""Сборка единственного листа «Данные» из книги ГТ и выгрузки датчиков.

Excel здесь является источником измерений и профиля поправок. Расчёты
выполняются над таблицей в памяти, поэтому позднее источник можно заменить БД.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import json
import os
import re

import numpy as np
import pandas as pd
from CoolProp.CoolProp import PropsSI
from CoolProp.HumidAirProp import HAPropsSI
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import column_index_from_string, get_column_letter
from station_profiles import (
    StationRepository, get_station_repository, resolve_station,
    sensor_target_columns, unit_number,
)
SERVICE_FIELDS = {"date", "time", "milliseconds"}


@dataclass(frozen=True)
class SensorMapping:
    source: str
    target: str
    factor: float = 1.0


@dataclass
class AssembledData:
    values: pd.DataFrame  # Внутренние имена колонок — Excel-буквы.
    headers: dict[str, str | None]
    imported_rows: int
    overlap_rows: int
    invalid_sensor_values: dict[str, int]
    verified_duplicate_columns: list[str]
    undated_rows: pd.DataFrame | None = None  # Сохраняются при экспорте, в расчёт не входят.

    def as_headered_frame(self) -> pd.DataFrame:
        result = self.values.reset_index(drop=True).copy()
        result.columns = [
            self.headers.get(column) or f"Unnamed: {index}"
            for index, column in enumerate(result.columns)
        ]
        return result


def list_source_sheets(path: str | Path) -> list[str]:
    with pd.ExcelFile(path, engine="calamine") as workbook:
        return list(workbook.sheet_names)


def _header_row(workbook: pd.ExcelFile, sheet_name: str) -> int:
    head = workbook.parse(sheet_name, header=None, nrows=60)
    for index, row in head.iterrows():
        labels = {str(value).strip().lower() for value in row if pd.notna(value)}
        if {"date", "time"} <= labels:
            return int(index)
    raise ValueError(f"На листе «{sheet_name}» не найдены колонки date и time.")


def source_columns(path: str | Path, sheet_name: str) -> list[str]:
    with pd.ExcelFile(path, engine="calamine") as workbook:
        row = _header_row(workbook, sheet_name)
        head = workbook.parse(sheet_name, header=None, nrows=row + 1).iloc[row]
    return [
        str(value).strip() for value in head
        if pd.notna(value) and str(value).strip()
    ]


def template_columns(path: str | Path) -> dict[str, str | None]:
    with pd.ExcelFile(path, engine="calamine") as workbook:
        if not workbook.sheet_names or workbook.sheet_names[0] != "Данные":
            raise ValueError("Первый лист расчётной книги должен называться «Данные».")
        raw = workbook.parse("Данные", header=None)
        row = raw.iloc[0]
    headed = [index for index, value in enumerate(row) if pd.notna(value)]
    if not headed:
        raise ValueError("На листе «Данные» нет заголовков колонок.")
    return {
        get_column_letter(index + 1): None if pd.isna(value) else str(value)
        for index, value in enumerate(row)
    }


def _mapping_profile_path() -> Path:
    base = Path(os.environ.get("LOCALAPPDATA") or Path.home() / ".config")
    return base / "GTU_Report" / "sensor_mappings.json"


def _profile_key(book_path: str | Path) -> str:
    stem = Path(book_path).stem
    match = re.search(
        r"^(.*?)[\s_-]+(?:ГТ|GT)[-\s_]*(\d+)(?=$|[^0-9])",
        stem, re.I,
    )
    if not match:
        raise ValueError(
            "В имени книги не найдены станция и номер ГТ для профиля поправок."
        )
    station = re.sub(r"[\s_]+", " ", match.group(1).strip()).casefold()
    return f"{station}|{match.group(2)}"


def load_saved_mappings(
    book_path: str | Path, sheet_headers: list[str],
    *, station_id: str | None = None,
    repository: StationRepository | None = None,
) -> list[SensorMapping] | None:
    station = resolve_station(book_path, station_id, repository)
    if station is None:
        return None
    path = _mapping_profile_path()
    if not path.exists():
        return None
    try:
        saved = json.loads(path.read_text(encoding="utf-8"))
        key = f"{station.id}|{unit_number(book_path) or ''}|{'|'.join(sheet_headers)}"
        entries = saved.get(key)
        # Старые назначения переносим только при совпадении станции по имени.
        detected = resolve_station(book_path, repository=repository)
        if entries is None and detected is not None and detected.id == station.id:
            entries = saved.get(f"{_profile_key(book_path)}|{'|'.join(sheet_headers)}")
        if entries is None:
            return None
        return [SensorMapping(**entry) for entry in entries]
    except (OSError, ValueError, TypeError, KeyError):
        return None


def save_mappings(
    book_path: str | Path, sheet_headers: list[str],
    mappings: list[SensorMapping],
    *, station_id: str | None = None,
    repository: StationRepository | None = None,
) -> None:
    station = resolve_station(book_path, station_id, repository)
    if station is None:
        raise ValueError("Выберите ТЭЦ перед сохранением соответствий датчиков.")
    path = _mapping_profile_path()
    try:
        saved = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
        if not isinstance(saved, dict):
            saved = {}
    except (OSError, ValueError):
        saved = {}
    key = f"{station.id}|{unit_number(book_path) or ''}|{'|'.join(sheet_headers)}"
    saved[key] = [
        {"source": item.source, "target": item.target, "factor": item.factor}
        for item in mappings
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(saved, ensure_ascii=False, indent=2), encoding="utf-8")


def suggest_mappings(
    source_headers: list[str], book_path: str | Path,
    verified_duplicates: set[str] | None = None,
    *, station_id: str | None = None,
    repository: StationRepository | None = None,
) -> list[SensorMapping]:
    """Предлагает только каналы выбранной ТЭЦ и ГТ; неизвестные — вручную."""
    repository = repository or get_station_repository()
    station = resolve_station(book_path, station_id, repository)
    gt_number = unit_number(book_path)
    channels = repository.get_channels(station.id) if station else ()
    by_code = {channel.code.casefold(): channel for channel in channels
               if channel.unit_number == gt_number}
    suggestions: list[SensorMapping] = []
    for header in source_headers:
        if header.lower() in SERVICE_FIELDS or header in (verified_duplicates or set()):
            continue
        channel = by_code.get(header.casefold())
        suggestions.append(SensorMapping(header, channel.target_column, channel.factor)
                           if channel else SensorMapping(header, "", 1.0))
    return suggestions


def validate_mappings(
    mappings: list[SensorMapping], source_headers: list[str],
    verified_duplicates: set[str] | None = None,
    target_headers: dict[str, str | None] | None = None,
    *, allowed_targets: set[str] | None = None,
) -> None:
    allowed_targets = sensor_target_columns() if allowed_targets is None else allowed_targets
    available = set(source_headers)
    seen_sources: set[str] = set()
    seen_targets: set[str] = set()
    for mapping in mappings:
        if mapping.source not in available:
            raise ValueError(f"В выгрузке нет канала {mapping.source!r}.")
        if mapping.source in seen_sources:
            raise ValueError(f"Канал {mapping.source!r} назначен дважды.")
        if not mapping.target or not re.fullmatch(r"[A-Z]{1,3}", mapping.target):
            raise ValueError(f"Для канала {mapping.source!r} не выбрана колонка листа.")
        if mapping.target in {"A", "B", "AD"}:
            raise ValueError(f"Колонка {mapping.target} зарезервирована.")
        if target_headers is not None and mapping.target not in allowed_targets:
            last = max(column_index_from_string(col) for col in target_headers)
            if column_index_from_string(mapping.target) <= last:
                raise ValueError(
                    f"Колонка {mapping.target} относится к расчётной части книги. "
                    f"Для нового канала используйте колонку после {get_column_letter(last)}."
                )
        if mapping.target in seen_targets:
            raise ValueError(f"На колонку {mapping.target} назначено два канала.")
        if not np.isfinite(mapping.factor) or mapping.factor == 0:
            raise ValueError(f"Неверный коэффициент для {mapping.source!r}.")
        seen_sources.add(mapping.source)
        seen_targets.add(mapping.target)
    missing = {
        header for header in source_headers
        if header.lower() not in SERVICE_FIELDS
        and header not in (verified_duplicates or set())
        and header not in seen_sources
    }
    if missing:
        raise ValueError(f"Не назначены каналы: {', '.join(sorted(missing))}.")


def _verified_duplicates(raw: pd.DataFrame, station_id: str,
                         repository: StationRepository | None = None) -> set[str]:
    """Разрешает пропуск только точно совпавших производных колонок."""
    repository = repository or get_station_repository()
    result = set()
    for relation in repository.get_duplicates(station_id):
        derived, original = relation.derived_code, relation.source_code
        if derived not in raw or original not in raw:
            continue
        left = pd.to_numeric(raw[derived], errors="coerce")
        right = pd.to_numeric(raw[original], errors="coerce") * relation.factor
        both = left.notna() & right.notna()
        if both.any() and (left.notna() == right.notna()).all() and np.allclose(
            left[both], right[both], rtol=0, atol=1e-8
        ):
            result.add(derived)
    return result


def _read_sensor_sheet(path: str | Path, sheet_name: str) -> pd.DataFrame:
    with pd.ExcelFile(path, engine="calamine") as workbook:
        row = _header_row(workbook, sheet_name)
        raw = workbook.parse(sheet_name, header=row)
    raw = raw.loc[:, ~raw.columns.astype(str).str.startswith("Unnamed:")].copy()
    raw.columns = [str(column).strip() for column in raw.columns]
    date = pd.to_datetime(raw["date"], errors="coerce")
    time = pd.to_timedelta(raw["time"].astype(str), errors="coerce")
    timestamp = date + time
    valid = timestamp.notna() & (timestamp.dt.minute % 30 == 0) & (
        timestamp.dt.second == 0
    )
    selected = raw.loc[valid].copy()
    selected.index = pd.DatetimeIndex(timestamp.loc[valid], name="timestamp")
    if selected.index.has_duplicates:
        raise ValueError("В выбранном листе повторяются даты и время замеров.")
    if selected.empty:
        raise ValueError("На выбранном листе нет получасовых замеров.")
    return selected.sort_index()


def source_mapping_info(
    path: str | Path, sheet_name: str,
    *, station_id: str,
    repository: StationRepository | None = None,
) -> tuple[list[str], set[str]]:
    """Каналы листа и производные колонки, проверенные по числам."""
    raw = _read_sensor_sheet(path, sheet_name)
    return list(raw.columns), _verified_duplicates(raw, station_id, repository)


def _read_template(path: str | Path) -> tuple[pd.DataFrame, dict[str, str | None], pd.DataFrame]:
    with pd.ExcelFile(path, engine="calamine") as workbook:
        if not workbook.sheet_names or workbook.sheet_names[0] != "Данные":
            raise ValueError("Первый лист расчётной книги должен называться «Данные».")
        raw = workbook.parse("Данные", header=None)
    headed = [index for index, value in enumerate(raw.iloc[0]) if pd.notna(value)]
    if not headed:
        raise ValueError("На листе «Данные» нет заголовков колонок.")
    headers = {
        get_column_letter(index + 1): None if pd.isna(value) else str(value)
        for index, value in enumerate(raw.iloc[0])
    }
    data = raw.iloc[1:].copy().astype(object)
    data.columns = list(headers)
    dates = pd.to_datetime(data["B"], errors="coerce")
    undated_rows = data.loc[dates.isna()].copy()
    rounded = dates.dt.round("min")
    excessive_shift = dates.notna() & (
        (dates - rounded).abs() > pd.Timedelta(seconds=1)
    )
    if excessive_shift.any():
        raise ValueError(
            "В расчётной книге найдены даты со смещением больше секунды "
            "от целой минуты; автоматическое совмещение небезопасно."
        )
    dates = rounded
    data = data.loc[dates.notna()].copy()
    data.index = pd.DatetimeIndex(dates.loc[dates.notna()], name="timestamp")
    if data.index.has_duplicates:
        raise ValueError("В расчётной книге повторяются даты на листе «Данные».")
    data["B"] = data.index
    return data.sort_index(), headers, undated_rows


def _convert_sensor(series: pd.Series, factor: float) -> tuple[pd.Series, int]:
    values = pd.to_numeric(series, errors="coerce")
    invalid = int((series.notna() & values.isna()).sum())
    converted = pd.Series(values * factor, index=series.index, dtype=object)
    converted.loc[series.notna() & values.isna()] = series.loc[
        series.notna() & values.isna()
    ]
    return converted, invalid


def assemble_sensor_data(
    book_path: str | Path,
    source_path: str | Path,
    source_sheet: str,
    mappings: list[SensorMapping],
    *, station_id: str | None = None,
    repository: StationRepository | None = None,
) -> AssembledData:
    repository = repository or get_station_repository()
    station = resolve_station(book_path, station_id, repository)
    if station is None:
        raise ValueError("Выберите ТЭЦ перед обновлением данных датчиков.")
    template, headers, undated_rows = _read_template(book_path)
    raw = _read_sensor_sheet(source_path, source_sheet)
    duplicates = _verified_duplicates(raw, station.id, repository)
    validate_mappings(mappings, list(raw.columns), duplicates, headers,
                      allowed_targets=sensor_target_columns(station.id, repository))
    gt_number = unit_number(book_path)
    if gt_number:
        channels = repository.get_channels(station.id)
        codes = {str(col).casefold() for col in raw.columns}
        expected_count = sum(channel.code.casefold() in codes for channel in channels
                             if channel.unit_number == gt_number)
        other_count = sum(channel.code.casefold() in codes for channel in channels
                          if channel.unit_number != gt_number)
        if other_count > expected_count:
            raise ValueError(
                f"Книга выбрана для ГТ-{gt_number}, но каналы листа "
                f"«{source_sheet}» относятся к другому агрегату."
            )

    all_dates = template.index.union(raw.index).sort_values()
    result = template.reindex(all_dates).copy()
    # В старой книге давление могло лежать в AD вместо AC. Заполняем
    # только пропуск AC, сохраняя обе исходные колонки без перезаписи.
    if "AD" in result and "AC" in result and not headers.get("AD"):
        shifted = template["AC"].isna() & template["AD"].notna()
        result.loc[template.index[shifted], "AC"] = template.loc[
            shifted, "AD"
        ].to_numpy()
    invalid_values: dict[str, int] = {}
    for mapping in mappings:
        converted, invalid = _convert_sensor(raw[mapping.source], mapping.factor)
        if mapping.target not in result:
            result[mapping.target] = np.nan
            headers[mapping.target] = mapping.source
        result.loc[raw.index, mapping.target] = converted.to_numpy()
        if invalid:
            invalid_values[mapping.source] = invalid

    result["B"] = result.index
    result["A"] = np.arange(len(result), dtype=float) * 0.5
    return AssembledData(
        values=result,
        headers=headers,
        imported_rows=len(raw),
        overlap_rows=len(template.index.intersection(raw.index)),
        invalid_sensor_values=invalid_values,
        verified_duplicate_columns=sorted(duplicates),
        undated_rows=undated_rows,
    )


def _profile_cells() -> set[str]:
    groups = (
        (669, "MN"), (787, "TVX"), (788, ("AB", "AD", "AF")),
        (888, ("V", "W", "X", "AB", "AC", "AD")),
        (712, "MNOP"), (507, "PQRST"),
        (969, ("T", "U", "V", "AC", "AD", "AE")),
        (1107, "TU"),
        (1019, ("L", "M", "N", "R", "S", "T")),
        (851, (
            "U", "V", "W", "X", "Y", "Z", "AC", "AD", "AE", "AF",
            "AG", "AH", "AK", "AL", "AM", "AN", "AO", "AP", "AS",
            "AT", "AU", "AV", "AW", "AX",
        )),
        (605, "MO"),
    )
    return {f"{column}{row}" for row, columns in groups for column in columns}


def _calculation_profile_path() -> Path:
    return _mapping_profile_path().with_name("calculation_profiles.json")


def _validate_calculation_profile(profile: dict) -> dict[str, float]:
    required = _profile_cells() | {
        "gas_c_h", "gas_density", "qnr_denominator"
    }
    try:
        numbers = {key: float(profile[key]) for key in required}
    except (KeyError, TypeError, ValueError) as error:
        raise ValueError("Сохранённый профиль поправок неполон.") from error
    if not all(np.isfinite(value) for value in numbers.values()):
        raise ValueError("В профиле поправок есть нечисловые значения.")
    if numbers["gas_density"] <= 0 or numbers["qnr_denominator"] <= 0:
        raise ValueError("В профиле поправок некорректная плотность газа.")
    if "gas_q_lower" in profile:
        q_lower = float(profile["gas_q_lower"])
        if not np.isfinite(q_lower) or q_lower <= 0:
            raise ValueError("В паспорте газа некорректное Qнр.")
        numbers["gas_q_lower"] = q_lower
    return numbers


def _load_cached_calculation_profile(book_path: str | Path) -> dict[str, float]:
    path = _calculation_profile_path()
    try:
        saved = json.loads(path.read_text(encoding="utf-8"))
        profile = saved[_profile_key(book_path)]
    except (OSError, KeyError, TypeError, ValueError) as error:
        raise ValueError(
            "В выбранной книге нет листа «Поправки», а прежний локальный профиль "
            "для этой станции и ГТ недоступен. Выберите исходную расчётную книгу "
            "с листом «Поправки». Паспорт газа задаётся отдельно."
        ) from error
    print("Использован сохранённый профиль поправок для этой станции и ГТ.")
    return _validate_calculation_profile(profile)


def _save_calculation_profile(
    book_path: str | Path, profile: dict[str, float]
) -> None:
    path = _calculation_profile_path()
    try:
        saved = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
        if not isinstance(saved, dict):
            saved = {}
    except (OSError, ValueError):
        saved = {}
    saved[_profile_key(book_path)] = _validate_calculation_profile(profile)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(saved, ensure_ascii=False, indent=2), encoding="utf-8")


def read_calculation_profile(book_path: str | Path, *, gas_passport=None) -> dict[str, float]:
    """Читает только используемые поправки и постоянные выбранной ГТ."""
    with pd.ExcelFile(book_path, engine="calamine") as workbook:
        if "Поправки" not in workbook.sheet_names:
            profile = _load_cached_calculation_profile(book_path)
            return _with_gas_passport(profile, gas_passport)
        corrections = workbook.parse("Поправки", header=None)
        results = (workbook.parse("Результаты", header=None, nrows=2)
                   if "Результаты" in workbook.sheet_names else None)
    profile: dict[str, float] = {}
    for cell in _profile_cells():
        match = re.fullmatch(r"([A-Z]+)(\d+)", cell)
        assert match is not None
        from openpyxl.utils import column_index_from_string
        column, row = match.groups()
        try:
            value = corrections.iat[int(row) - 1, column_index_from_string(column) - 1]
        except IndexError as error:
            raise ValueError(f"На листе «Поправки» нет ячейки {cell} принятой структуры СызТЭЦ.") from error
        try:
            profile[cell] = float(value)
        except (TypeError, ValueError) as error:
            raise ValueError(f"Нет числовой поправки {cell} для этой ГТ.") from error
    if gas_passport is not None:
        return _with_gas_passport(profile, gas_passport)
    # Совместимость с прямым старым вызовом: экспортированная книга имеет
    # паспорт на «Поправки» и не требует прежнего листа «Результаты».
    from gas_passport import embedded_passport
    from station_profiles import unit_number
    station = resolve_station(book_path)
    turbine = unit_number(book_path)
    if station is not None and turbine:
        embedded = embedded_passport(corrections, station.id, turbine)
        if embedded is not None:
            return _with_gas_passport(profile, embedded)
    if results is None:
        raise ValueError("Нет паспорта газа. Заполните его кнопкой «Паспорт газа» для выбранной ТЭЦ и ГТ.")
    profile["gas_c_h"] = float(results.iat[1, 30])  # Совместимость прежнего API.
    profile["gas_density"] = float(results.iat[1, 28])
    # В ГТ-10 AK2 делит на Результаты!AC2, а в ГТ-11 — на 0.698.
    formula_book = load_workbook(book_path, read_only=True, data_only=False)
    try:
        ak_formula = str(formula_book["Данные"]["AK2"].value or "")
    finally:
        formula_book.close()
    direct_divisor = re.fullmatch(r"=34\.26/([0-9.]+)", ak_formula)
    profile["qnr_denominator"] = (
        float(direct_divisor.group(1)) if direct_divisor
        else profile["gas_density"]
    )
    profile = _validate_calculation_profile(profile)
    try:
        _save_calculation_profile(book_path, profile)
    except (OSError, ValueError) as error:
        print(f"Не удалось сохранить профиль поправок: {error}")
    return profile


def _with_gas_passport(profile, passport):
    if passport is None:
        return _validate_calculation_profile(profile)
    profile = {**profile, "gas_c_h": passport.carbon_hydrogen_ratio,
               "gas_density": passport.density_kg_per_nm3,
               "qnr_denominator": passport.density_kg_per_nm3,
               "gas_q_lower": passport.q_lower_mj_per_nm3}
    return _validate_calculation_profile(profile)


def _number(frame: pd.DataFrame, column: str) -> np.ndarray:
    if column not in frame:
        return np.full(len(frame), np.nan, dtype=float)
    return pd.to_numeric(frame[column], errors="coerce").to_numpy(dtype=float)


def _polynomial(profile: dict[str, float], row: int, columns: tuple[str, ...],
                x: np.ndarray) -> np.ndarray:
    return np.polyval([profile[f"{column}{row}"] for column in columns], x)


def _coolprop_columns(frame: pd.DataFrame) -> None:
    """Влажность, плотность и термодинамика по формулам книги."""
    c, f, g = (_number(frame, col) for col in ("C", "F", "G"))
    n, p, w, x = (_number(frame, col) for col in ("N", "P", "W", "X"))
    count = len(frame)
    result = {col: np.full(count, np.nan, dtype=float) for col in (
        "I", "Q", "R", "Z", "AA", "AB"
    )}
    reference_h = PropsSI("H", "T", 273.15, "P", 101325, "HEOS::Air")
    for i in range(count):
        if np.isfinite(c[i]) and c[i] > 0 and np.isfinite(f[i]) and np.isfinite(g[i]):
            try:
                result["I"][i] = HAPropsSI(
                    "W", "T", f[i] + 273.15, "P", c[i],
                    "R", max(g[i], 0.0) / 100.0,
                )
            except (ValueError, OverflowError):
                pass
        if np.isfinite(n[i]) and n[i] > 0 and np.isfinite(p[i]):
            if np.isfinite(result["I"][i]):
                try:
                    result["Q"][i] = 100.0 * HAPropsSI(
                        "R", "T", p[i] + 273.15, "P", n[i],
                        "W", result["I"][i],
                    )
                    result["R"][i] = 1.0 / HAPropsSI(
                        "Vha", "T", p[i] + 273.15, "P", n[i],
                        "R", result["Q"][i] / 100.0,
                    )
                except (ValueError, OverflowError, ZeroDivisionError):
                    pass
            try:
                h1 = PropsSI("H", "T", p[i] + 273.15, "P", n[i], "HEOS::Air")
                result["Z"][i] = (h1 - reference_h) / 1000.0
                if np.isfinite(x[i]) and x[i] > 0 and np.isfinite(w[i]):
                    outlet_p = x[i] * 1_000_000.0
                    h2 = PropsSI(
                        "H", "T", w[i] + 273.15, "P", outlet_p, "HEOS::Air"
                    )
                    result["AA"][i] = (h2 - reference_h) / 1000.0
                    if outlet_p > n[i] and abs(h2 - h1) > 1e-10:
                        s1 = PropsSI(
                            "S", "T", p[i] + 273.15, "P", n[i], "HEOS::Air"
                        )
                        h2s = PropsSI("H", "S", s1, "P", outlet_p, "HEOS::Air")
                        result["AB"][i] = (h2s - h1) / (h2 - h1) * 100.0
            except (ValueError, OverflowError):
                pass
    for column, values in result.items():
        frame[column] = values


def calculate_data_columns(data: AssembledData,
                           profile: dict[str, float]) -> AssembledData:
    """Пересчитывает производные колонки без ссылок на другие листы."""
    frame = data.values
    q = lambda cell: profile[cell]
    c, f, g, l, m, p, x = (
        _number(frame, col) for col in ("C", "F", "G", "L", "M", "P", "X")
    )
    frame["D"] = q("M669") * c / 1000.0 + q("N669")
    frame["H"] = (
        (q("T787") * p**3 + q("V787") * p**2 + q("X787")) * g
        + q("AB788") * p**3 + q("AD788") * p**2 + q("AF788")
    )
    frame["N"] = c - m
    frame["O"] = (
        _polynomial(profile, 888, ("V", "W", "X"), p) * l
        + _polynomial(profile, 888, ("AB", "AC", "AD"), p)
    )
    frame["S"] = p - f
    frame["T"] = _polynomial(profile, 712, ("M", "N", "O", "P"), p)
    frame["U"] = _polynomial(profile, 507, ("P", "Q", "R", "S", "T"), p)
    with np.errstate(divide="ignore", invalid="ignore"):
        frame["Y"] = x * 1_000_000.0 / _number(frame, "N")
    frame["AE"] = (
        _polynomial(profile, 969, ("T", "U", "V"), p)
        * _number(frame, "AC")
        + _polynomial(profile, 969, ("AC", "AD", "AE"), p)
    )
    frame["AI"] = q("T1107") * _number(frame, "AH") + q("U1107")
    rpm = _number(frame, "AM") / 5200.0
    frame["AN"] = (
        _polynomial(profile, 851, ("U", "V", "W", "X", "Y", "Z"), p) * rpm**3
        + _polynomial(profile, 851, ("AC", "AD", "AE", "AF", "AG", "AH"), p) * rpm**2
        + _polynomial(profile, 851, ("AK", "AL", "AM", "AN", "AO", "AP"), p) * rpm
        + _polynomial(profile, 851, ("AS", "AT", "AU", "AV", "AW", "AX"), p)
    )
    _coolprop_columns(frame)
    recalculate_gas_columns(frame, profile)
    return data


def recalculate_gas_columns(frame: pd.DataFrame, profile: dict[str, float]) -> None:
    """Меняет только газозависимую часть, остальные поправки остаются прежними."""
    q = lambda cell: profile[cell]
    frame["AJ"] = profile["gas_c_h"]
    # Старый вызов без паспорта сохраняет прежнюю методику; GUI передаёт
    # паспортный gas_q_lower и плотность, без фиксированного числителя.
    frame["AK"] = profile.get("gas_q_lower", 34.26) / profile["qnr_denominator"]
    aj, ak = (_number(frame, col) for col in ("AJ", "AK"))
    frame["AL"] = (
        _polynomial(profile, 1019, ("L", "M", "N"), aj) * ak
        + _polynomial(profile, 1019, ("R", "S", "T"), aj)
    )
    denominator = np.ones(len(frame), dtype=float)
    for column in ("D", "H", "O", "T", "AE", "AI", "AL", "AN"):
        denominator *= _number(frame, column)
    with np.errstate(divide="ignore", invalid="ignore"):
        power = _number(frame, "AO") / denominator
    power[~np.isfinite(power)] = np.nan
    frame["AQ"] = power
    density = profile.get("gas_density")
    frame["AT"] = _number(frame, "AS") * 3600.0 / (density if "gas_q_lower" in profile else 0.698)
    frame["AV"] = (power * q("M605") + q("O605")) * 3600.0 / (density if "gas_q_lower" in profile else 0.6993)
    with np.errstate(divide="ignore", invalid="ignore"):
        specific = _number(frame, "AV") / power
    specific[~np.isfinite(specific)] = np.nan
    frame["AX"] = specific
    for output, source in (("AP", "AO"), ("AR", "AQ"),
                           ("AU", "AT"), ("AW", "AV")):
        values = _number(frame, source)
        averaged = np.full(len(frame), np.nan)
        for i in range(1, len(frame), 2):
            pair = values[i - 1:i + 1]
            if np.isfinite(pair).any():
                averaged[i] = float(np.nanmean(pair))
        frame[output] = averaged


def apply_economic_columns(data: AssembledData, economics_result: dict) -> None:
    """Записывает построчные потери действующей методики в колонки книги."""
    frame = data.values
    loss_columns = [
        column for column, header in data.headers.items()
        if header and header.strip().lower() in {"перерасход", "недовыработка"}
    ]
    for column in loss_columns:
        frame[column] = np.nan
        next_column = get_column_letter(column_index_from_string(column) + 1)
        if data.headers.get(next_column) in (None, ""):
            frame[next_column] = np.nan
            data.headers.setdefault(next_column, None)
    records = economics_result.get("row_losses")
    if records is None or records.empty:
        return
    losses = records.copy()
    losses["date"] = pd.to_datetime(losses["date"]).dt.round("min")
    if losses["date"].duplicated().any():
        raise ValueError("Экономический расчёт вернул повторяющиеся даты.")
    losses = losses.set_index("date").sort_index()
    source_for_header = {
        "перерасход": "fuel_loss_rate",
        "недовыработка": "power_loss_rate",
    }
    for column, header in list(data.headers.items()):
        if header is None:
            continue
        source = source_for_header.get(header.strip().lower())
        if source is None:
            continue
        rates = losses[source].reindex(frame.index).fillna(0.0)
        frame[column] = rates.to_numpy(dtype=float)
        next_column = get_column_letter(column_index_from_string(column) + 1)
        if data.headers.get(next_column) not in (None, ""):
            continue
        averaged = pd.Series(np.nan, index=frame.index, dtype=float)
        for _, group in losses.groupby("period", sort=False):
            dates = group.index.sort_values()
            for index in range(0, len(dates) - 1, 2):
                pair = rates.reindex(dates[index:index + 2]).to_numpy(dtype=float)
                averaged.loc[dates[index + 1]] = float(pair.mean())
        frame[next_column] = averaged.to_numpy()
        data.headers.setdefault(next_column, None)


def save_data_sheet(data: AssembledData, output_path: str | Path) -> Path:
    """Сохраняет только «Данные» как значения, без внешних ссылок и макросов."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Данные")
    sheet.freeze_panes = "B2"
    last_column = max(column_index_from_string(col) for col in data.values.columns)
    letters = [get_column_letter(index) for index in range(1, last_column + 1)]
    sheet.append([data.headers.get(col) for col in letters])
    from itertools import chain
    rows = chain(
        data.values.reindex(columns=letters).itertuples(index=False, name=None),
        data.undated_rows.reindex(columns=letters).itertuples(index=False, name=None)
        if data.undated_rows is not None else (),
    )
    for row in rows:
        output = []
        for value in row:
            if pd.isna(value):
                output.append(None)
            elif isinstance(value, pd.Timestamp):
                output.append(value.to_pydatetime())
            elif isinstance(value, np.generic):
                output.append(value.item())
            elif isinstance(value, str) and value.startswith("="):
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"
                output.append(cell)
            else:
                output.append(value)
        sheet.append(output)
    book.save(path)
    return path
