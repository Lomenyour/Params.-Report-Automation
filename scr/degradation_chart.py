"""Редактируемая диаграмма деградации на листе «Графики»."""
from __future__ import annotations

from dataclasses import dataclass

import numpy as np
import pandas as pd
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines, NumericAxis
from openpyxl.chart.data_source import NumData, NumVal
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties,
    RichTextProperties,
)
from openpyxl.utils.datetime import to_excel
from chart_text import FONT_SIZE, rich_text, label_series

from RH_raschet import _dominant_vna_interval


HELPER_COLUMN = 62  # BJ:BN; правее двух видимых диаграмм и их таблиц.
HELPER_HEADERS = ("Дата", "Nприв, МВт", "bприв", "ηГТУ, %", "ВНА, %")
METRICS = (
    ("power", "Nприв, МВт", "C00000"),
    ("fuel_specific", "bприв", "ED7D31"),
    ("eta_gtu", "ηГТУ, %", "7030A0"),
    ("vna", "ВНА, %", "00A6D6"),
)


@dataclass
class DegradationPlot:
    points: pd.DataFrame
    periods: list[dict]


def prepare_degradation_plot(
    values: pd.DataFrame | None,
    periods: pd.DataFrame | None,
    delta_table: pd.DataFrame | None,
) -> DegradationPlot | None:
    """Выбирает те же одноградусные интервалы ВНА, что и регрессия."""
    if values is None or values.empty or periods is None or periods.empty:
        return None

    saved_intervals = (
        delta_table.attrs.get("vna_intervals", {})
        if delta_table is not None else {}
    )
    selected_parts = []
    descriptions = []

    for row in periods.itertuples(index=False):
        number = int(row.period)
        if "period" in values.columns:
            period_data = values.loc[values["period"] == number]
        else:
            dates = pd.to_datetime(values["date"], errors="coerce")
            period_data = values.loc[
                (dates >= row.date_start) & (dates <= row.date_end)
            ]
        if period_data.empty:
            continue

        interval = saved_intervals.get(number)
        if interval is None:
            interval = _dominant_vna_interval(period_data)
        if interval is None:
            continue

        lower, upper, _ = interval
        vna = pd.to_numeric(period_data["vna"], errors="coerce")
        selected = period_data.loc[(vna >= lower) & (vna < upper)].copy()
        selected["date"] = pd.to_datetime(selected["date"], errors="coerce")
        selected = selected.loc[selected["date"].notna()]
        if selected.empty:
            continue
        for key, _, _ in METRICS:
            if key not in selected.columns:
                selected[key] = np.nan
            selected[key] = pd.to_numeric(selected[key], errors="coerce")
            selected.loc[~np.isfinite(selected[key]), key] = np.nan
        selected.loc[selected["eta_gtu"] <= 0, "eta_gtu"] = np.nan
        selected_parts.append(
            selected[["date", "power", "fuel_specific", "eta_gtu", "vna"]]
        )

        delta = None
        if delta_table is not None and not delta_table.empty:
            match = delta_table.loc[delta_table["Период"] == number]
            if not match.empty:
                delta = match.iloc[0]
        descriptions.append({
            "period": number,
            "date_start": pd.Timestamp(row.date_start),
            "date_end": pd.Timestamp(row.date_end),
            "vna_start": float(lower),
            "vna_end": float(upper),
            "count": len(selected),
            "points_start": selected["date"].min(),
            "points_end": selected["date"].max(),
            "levels": {key: _finite(selected[key].median()) for key, _, _ in METRICS},
            "delta_n": _finite(delta.get("ΔNприв")) if delta is not None else None,
            "delta_b": _finite(delta.get("Δbприв")) if delta is not None else None,
            "delta_eta": _finite(delta.get("Δη, п.п.")) if delta is not None else None,
        })

    if not selected_parts:
        return None
    points = pd.concat(selected_parts, ignore_index=True)
    points = points.sort_values("date", kind="stable").reset_index(drop=True)
    return DegradationPlot(points, descriptions)


def _finite(value):
    parsed = pd.to_numeric(value, errors="coerce")
    return float(parsed) if pd.notna(parsed) and np.isfinite(parsed) else None


def _cache(values, *, dates=False):
    points = []
    for index, value in enumerate(values):
        if pd.isna(value):
            continue
        number = (
            to_excel(pd.Timestamp(value).to_pydatetime())
            if dates else _finite(value)
        )
        if number is not None:
            points.append(NumVal(idx=index, v=float(number)))
    return NumData(
        formatCode="dd.mm.yyyy" if dates else "General",
        ptCount=len(values), pt=points,
    )


def _series(sheet, plot: DegradationPlot, key: str, title: str,
            color: str, column: int) -> Series:
    last = len(plot.points) + 1
    x = Reference(sheet, min_col=HELPER_COLUMN, min_row=2, max_row=last)
    y = Reference(sheet, min_col=column, min_row=2, max_row=last)
    series = Series(y, x, title=title)
    series.xVal.numRef.numCache = _cache(plot.points["date"], dates=True)
    series.yVal.numRef.numCache = _cache(plot.points[key])
    series.graphicalProperties.line.noFill = True
    series.marker.symbol = "circle"
    series.marker.size = 2
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color
    series.smooth = False
    return series


def _text(size=FONT_SIZE, bold=False):
    return RichText(
        bodyPr=RichTextProperties(lIns=45000, rIns=45000, tIns=35000, bIns=35000),
        p=[Paragraph(r=[], pPr=ParagraphProperties(
            defRPr=CharacterProperties(
                sz=size, b=bold, solidFill="000000",
                latin=DrawingFont(typeface="Times New Roman"),
            ),
        ), endParaRPr=CharacterProperties(lang="ru-RU"))],
    )


def _grid(color, width):
    return ChartLines(spPr=GraphicalProperties(
        ln=LineProperties(solidFill=color, w=width),
    ))


def _axis_title(axis, title):
    axis.title = title
    axis.title.txPr = _text(bold=True)
    axis.title.tx.rich = rich_text(title, bold=True)


def _summary_cells(sheet, plot: DegradationPlot, start_row: int, start_col: int):
    """Подписи интервалов и величины для визуальной проверки графика."""
    offset = start_col - 1
    cells = {start_row: {start_col: (
        "ДЕГРАДАЦИЯ ГТУ — точки в самом частом интервале ВНА каждого периода",
        "period",
    )}}
    sheet.merged_cells.add(
        f"{_col(start_col)}{start_row}:{_col(start_col + 22)}{start_row}"
    )
    sheet.row_dimensions[start_row].height = 27

    header_row = start_row + 38
    layout = (
        (1, 3, "Период"),
        (4, 7, "Даты периода"),
        (8, 10, "ВНА-фильтр"),
        (11, 12, "Точек"),
        (13, 14, "ΔN, МВт"),
        (15, 16, "Потеря N"),
        (17, 18, "Δbприв"),
        (19, 20, "Рост b"),
        (21, 23, "Δη, п.п."),
    )
    cells[header_row] = {}
    for first, last, title in layout:
        sheet.merged_cells.add(
            f"{_col(first + offset)}{header_row}:{_col(last + offset)}{header_row}"
        )
        cells[header_row][first + offset] = (title, "header")
    sheet.row_dimensions[header_row].height = 24

    for index, period in enumerate(plot.periods, start=1):
        row = header_row + index
        delta_n = period["delta_n"]
        delta_b = period["delta_b"]
        data = (
            f"Период {period['period']}",
            f"{period['date_start']:%d.%m.%Y} – {period['date_end']:%d.%m.%Y}",
            f"{period['vna_start']:.0f}–{period['vna_end']:.0f} %",
            period["count"],
            delta_n,
            max(-delta_n, 0.0) if delta_n is not None else None,
            delta_b,
            max(delta_b, 0.0) if delta_b is not None else None,
            period["delta_eta"],
        )
        cells[row] = {}
        for (first, last, _), value in zip(layout, data):
            sheet.merged_cells.add(
                f"{_col(first + offset)}{row}:{_col(last + offset)}{row}"
            )
            cells[row][first + offset] = (value, "body")
        sheet.row_dimensions[row].height = 21
    return cells, header_row + len(plot.periods) + 1


def _col(index: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(index)


def _delta_label(symbol, value, unit):
    if value is None:
        return f"{symbol}: нет данных"
    if value == 0:
        return f"{symbol}: отсутствует"
    number = f"{abs(value):.3f}".rstrip("0").rstrip(".").replace(".", ",")
    return f"{symbol} = {'↑' if value > 0 else '↓'}{number} {unit}"


def _add_period_labels(chart, plot):
    first_label = len(chart.series)
    placed = {key: [] for key in ("fuel_specific", "vna", "power", "eta_gtu")}
    span = chart.x_axis.scaling.max - chart.x_axis.scaling.min
    for period in plot.periods:
        first = period["points_start"]
        moment = first + (period["points_end"] - first) / 2
        texts = (
            ("fuel_specific", _delta_label("Δbприв", period["delta_b"], "нм³/(МВт·ч)")),
            ("vna", f"{period['vna_start']:.1f} % ≤ ВНА < {period['vna_end']:.1f} %".replace(".", ",")),
            ("power", _delta_label("ΔNприв", period["delta_n"], "МВт")),
            ("eta_gtu", _delta_label("ΔηГТУ", period["delta_eta"], "п.п.")),
        )
        for key, text in texts:
            value = period["levels"][key]
            if value is None:
                continue
            y = (value - 65.) * 350. / 35. if key == "vna" else value
            offset = {"fuel_specific": 30., "vna": 25., "power": 25., "eta_gtu": 20.}[key]
            base = min(340., max(20., y + offset))
            x = to_excel(moment.to_pydatetime())
            # Короткие соседние периоды: оставляем X над своими данными,
            # но раздвигаем подписи по высоте, не сдвигая сами точки.
            nearby = [py for px, py in placed[key] if abs(x - px) < .20 * span]
            candidates = sorted({base + delta for delta in (0, 20, 40, 60, -20, -40, -60)}
                | {py + delta for py in nearby for delta in (20, -20)},
                key=lambda candidate: (candidate < base, abs(candidate - base)))
            y = next((candidate for candidate in candidates if 20 <= candidate <= 340
                and all(abs(x - px) >= .20 * span or abs(candidate - py) >= 18
                        for px, py in placed[key])), base)
            placed[key].append((x, y))
            chart.series.append(label_series(moment, y, text, dates=True, size=FONT_SIZE, bold=True))
    chart.legend.legendEntry = [LegendEntry(idx=index, delete=True)
                                for index in range(first_label, len(chart.series))]


def add_degradation_chart(sheet, plot: DegradationPlot,
                          start_row: int, start_col: int = 1):
    """Добавляет нативную XY-диаграмму и возвращает видимые ячейки таблицы."""
    main = ScatterChart()
    main.title = None
    main.roundedCorners = False
    main.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(noFill=True),
    )
    main.scatterStyle = "marker"
    main.width = 42
    main.height = 18
    main.x_axis.axId = 110
    main.y_axis.axId = 120
    main.x_axis.crossAx = 120
    main.y_axis.crossAx = 110
    main.x_axis.delete = False
    main.y_axis.delete = False
    _axis_title(main.x_axis, "Дата")
    main.x_axis.numFmt = "dd.mm.yyyy"
    main.x_axis.axPos = "b"
    main.x_axis.txPr = _text()
    _axis_title(main.y_axis, "Nприв, МВт; bприв; ηГТУ, %")
    main.y_axis.axPos = "l"
    main.y_axis.txPr = _text()
    main.y_axis.scaling.min = 0
    main.y_axis.scaling.max = 350
    main.y_axis.majorUnit = 50
    main.y_axis.minorUnit = 10
    main.y_axis.numFmt = "0"
    main.y_axis.tickLblPos = "nextTo"
    main.y_axis.majorGridlines = _grid("A6A6A6", 9525)
    main.y_axis.minorGridlines = _grid("D9D9D9", 4762)
    main.x_axis.majorGridlines = _grid("A6A6A6", 9525)
    main.x_axis.minorGridlines = _grid("D9D9D9", 4762)
    main.x_axis.tickLblPos = "low"
    main.y_axis.crosses = "min"
    main.legend.position = "b"
    main.legend.overlay = False
    main.legend.txPr = _text()
    main.legend.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.15, y=0.93, w=0.7, h=0.05,
    ))
    main.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.075, y=0.035, w=0.845, h=0.80,
    ))
    main.visible_cells_only = False
    main.display_blanks = "gap"

    first_date = plot.points["date"].min()
    last_date = plot.points["date"].max()
    lower_date = first_date.normalize() - pd.DateOffset(months=2)
    upper_date = last_date.normalize() + pd.DateOffset(months=2)
    main.x_axis.scaling.min = to_excel(lower_date.to_pydatetime())
    main.x_axis.scaling.max = to_excel(upper_date.to_pydatetime())
    span = max(1, (upper_date - lower_date).days)
    main.x_axis.majorUnit = next(
        (days for days in (1, 2, 5, 10, 20, 25, 50, 100, 150, 200, 365)
         if days >= span / 6), 365,
    )
    main.x_axis.minorUnit = main.x_axis.majorUnit / 5

    for offset, (key, title, color) in enumerate(METRICS[:3], start=1):
        main.series.append(_series(
            sheet, plot, key, title, color, HELPER_COLUMN + offset
        ))

    _add_period_labels(main, plot)

    secondary = ScatterChart()
    secondary.scatterStyle = "marker"
    secondary.x_axis = NumericAxis(
        axId=140, crossAx=130, axPos="b", delete=True,
    )
    secondary.x_axis.majorGridlines = None
    secondary.x_axis.minorGridlines = None
    secondary.x_axis.scaling.min = main.x_axis.scaling.min
    secondary.x_axis.scaling.max = main.x_axis.scaling.max
    secondary.y_axis.axId = 130
    secondary.y_axis.crossAx = 140
    secondary.y_axis.axPos = "r"
    secondary.y_axis.delete = False
    _axis_title(secondary.y_axis, "Положение ВНА, %")
    secondary.y_axis.txPr = _text()
    secondary.y_axis.crosses = "max"
    secondary.y_axis.tickLblPos = "nextTo"
    secondary.y_axis.numFmt = "0"
    secondary.y_axis.majorGridlines = None
    secondary.y_axis.minorGridlines = None
    secondary.y_axis.scaling.min = 65
    secondary.y_axis.scaling.max = 100
    secondary.y_axis.majorUnit = 5
    secondary.visible_cells_only = False
    secondary.series.append(_series(
        sheet, plot, "vna", "ВНА, %", "00A6D6", HELPER_COLUMN + 4
    ))
    main += secondary
    sheet.add_chart(main, f"{_col(start_col)}{start_row + 1}")
    return _summary_cells(sheet, plot, start_row, start_col)
