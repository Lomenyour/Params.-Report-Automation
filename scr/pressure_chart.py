"""Редактируемая Excel-диаграмма перепадов и таблицы граничных точек."""
from __future__ import annotations

from dataclasses import dataclass
import math

import numpy as np
import pandas as pd
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.axis import ChartLines, NumericAxis
from openpyxl.chart.data_source import NumData, NumVal
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineEndProperties, LineProperties
from openpyxl.drawing.text import (
    CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties,
    RichTextProperties,
)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.worksheet import Worksheet

from data_loader import SHORT_NAMES, _map_headers
from chart_events import ChartEvent
from degradation_chart import (
    HELPER_COLUMN as DEGRADATION_HELPER_COLUMN,
    HELPER_HEADERS as DEGRADATION_HELPER_HEADERS,
    add_degradation_chart,
    prepare_degradation_plot,
)


PRESSURE_CHANNELS = {
    "vlo": ("ВЛО, Па", "0070C0"),
    "fgo": ("ФГО, Па", "00B050"),
    "fto": ("ФТО, Па", "00B0F0"),
    "fgo_fto": ("ФГО+ФТО, Па", "00B050"),
}
POWER_COLOR = "C00000"
CHART_WIDTH_CM = 42
CHART_HEIGHT_CM = 18
TABLE_START_ROW = 40
HELPER_START_COLUMN = 50  # AX:BH; правее видимых диаграмм и таблиц.
PERIOD_HELPER_COLUMN = 59  # BG:BH: линии и подписи периодов.
DEGRADATION_CHART_COLUMN = 26  # Z: с зазором справа от диаграммы перепадов.
VISIBLE_LAST_COLUMN = DEGRADATION_CHART_COLUMN + 23  # AW: отступ для печати.


@dataclass
class PressureChartData:
    values: pd.DataFrame
    periods: pd.DataFrame
    boundaries: pd.DataFrame | None
    events: tuple[ChartEvent, ...] = ()
    degradation_values: pd.DataFrame | None = None
    degradation_deltas: pd.DataFrame | None = None


def _number(frame, key):
    if key not in frame:
        return pd.Series(np.nan, index=frame.index)
    return pd.to_numeric(frame[key], errors="coerce")


def pressure_channels(frame):
    """Совмещённый канал или отдельные ФГО и ФТО, только при наличии данных."""
    available = {
        key for key in PRESSURE_CHANNELS
        if np.isfinite(_number(frame, key)).any()
    }
    keys = ["vlo"] if "vlo" in available else []
    if "fgo_fto" in available:
        keys.append("fgo_fto")
    else:
        keys.extend(key for key in ("fgo", "fto") if key in available)
    return keys


def _source_columns(headers):
    mapped, _ = _map_headers(headers)
    result = {}
    for key, aliases in SHORT_NAMES.items():
        for alias in aliases:
            if alias in mapped:
                result[key] = mapped.index(alias) + 1
                break
    return result


def _text(size=1000, bold=False, centered=False):
    return RichText(
        bodyPr=RichTextProperties(lIns=45000, rIns=45000, tIns=35000, bIns=35000),
        p=[Paragraph(r=[], pPr=ParagraphProperties(
            algn="ctr" if centered else None,
            defRPr=CharacterProperties(
                sz=size, b=bold, solidFill="000000",
                latin=DrawingFont(typeface="Times New Roman"),
            ),
        ), endParaRPr=CharacterProperties(lang="ru-RU"))],
    )


def _grid(color, width):
    return ChartLines(spPr=GraphicalProperties(
        ln=LineProperties(solidFill=color, w=width),
    ))


def _axis_title(axis, title):
    axis.title = title
    axis.title.txPr = _text(1100, bold=True)
    for paragraph in axis.title.tx.rich.p:
        paragraph.pPr = ParagraphProperties(defRPr=CharacterProperties(
            sz=1100, b=True, latin=DrawingFont(typeface="Times New Roman"),
        ))
        for run in paragraph.r:
            run.rPr = CharacterProperties(
                sz=1100, b=True, latin=DrawingFont(typeface="Times New Roman"),
            )


def _scale_max(values):
    values = pd.to_numeric(pd.Series(values), errors="coerce")
    valid = values[np.isfinite(values) & (values >= 0)]
    maximum = float(valid.max()) if not valid.empty else 1.0
    if maximum <= 0:
        return 1.0
    step = 10 ** math.floor(math.log10(maximum)) / 5
    return math.ceil(maximum * 1.18 / step) * step


def _cache(values, *, dates=False):
    points = []
    for index, value in enumerate(values):
        if pd.isna(value):
            continue
        if dates:
            value = to_excel(pd.Timestamp(value).to_pydatetime())
        else:
            value = pd.to_numeric(value, errors="coerce")
            if pd.isna(value) or not np.isfinite(value) or value < 0:
                continue
        points.append(NumVal(idx=index, v=float(value)))
    return NumData(formatCode="dd-mm-yyyy" if dates else "General",
                   ptCount=len(values), pt=points)


def _point_series(sheet, key, title, color, frame, index):
    row_count = len(frame)
    x = Reference(sheet, min_col=HELPER_START_COLUMN, min_row=2, max_row=row_count + 1)
    y = Reference(sheet, min_col=HELPER_START_COLUMN + index, min_row=2, max_row=row_count + 1)
    series = Series(y, x, title=title)
    series.xVal.numRef.numCache = _cache(frame["date"], dates=True)
    series.yVal.numRef.numCache = _cache(_number(frame, key))
    series.graphicalProperties.line.noFill = True
    series.marker.symbol = "circle"
    series.marker.size = 2
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color
    series.smooth = False
    return series


def _period_series(sheet, row, points, title, label=False, arrow=False, vertical=False):
    count = len(points)
    x = Reference(sheet, min_col=PERIOD_HELPER_COLUMN, min_row=row, max_row=row + count - 1)
    y = Reference(sheet, min_col=PERIOD_HELPER_COLUMN + 1, min_row=row, max_row=row + count - 1)
    series = Series(y, x, title=title)
    series.xVal.numRef.numCache = _cache([point[0] for point in points], dates=True)
    series.yVal.numRef.numCache = _cache([point[1] for point in points])
    series.marker.symbol = "none"
    if label:
        series.graphicalProperties.line.noFill = True
        series.dLbls = DataLabelList(
            showSerName=True, showVal=False, showCatName=False,
            showLegendKey=False, dLblPos="ctr", txPr=_text(1000, centered=True),
            spPr=GraphicalProperties(
                solidFill="FFFFFF", ln=LineProperties(solidFill="000000", w=12700),
            ),
        )
        if vertical:
            series.dLbls.txPr.bodyPr.rot = -5400000
    else:
        series.graphicalProperties.line.solidFill = "000000"
        series.graphicalProperties.line.width = 15240
        if arrow:
            series.graphicalProperties.line.headEnd = LineEndProperties(type="triangle", w="sm", len="sm")
            series.graphicalProperties.line.tailEnd = LineEndProperties(type="triangle", w="sm", len="sm")
    return series


def _boundary_record(boundaries, number, tag):
    if boundaries is None or boundaries.empty:
        return None
    found = boundaries[boundaries["Период"].astype(str) == f"{number} {tag}"]
    return None if found.empty else found.iloc[0]


def _callout(record, keys):
    if record is None or str(record.get("дата", "—")) == "—":
        return "Нет номинальной точки\nВНА ≥ 80 %"
    date = _boundary_value(record.get("дата"), is_date=True)
    lines = [date.strftime("%d.%m.%Y") if isinstance(date, pd.Timestamp) else str(date)]

    def value(key):
        raw = record.get(key)
        parsed = pd.to_numeric(raw, errors="coerce")
        return f"{parsed:.1f}".replace(".", ",") if pd.notna(parsed) else "—"

    for key in keys:
        title = PRESSURE_CHANNELS[key][0].removesuffix(", Па")
        lines.append(f"ΔP {title} = {value(title)} Па;")
    lines += [f"ВНА = {value('ВНА')} %;", f"Nгту = {value('N')} МВт."]
    return "\n".join(lines)


def _callout_dates(periods, lower, upper):
    """Раздвигает близкие блоки, сохраняя порядок границ по времени."""
    dates = [pd.Timestamp(moment) for period in periods.itertuples(index=False)
             for moment in (period.date_start, period.date_end)]
    if not dates:
        return []
    span = upper - lower
    desired = np.array([(moment - lower) / span for moment in dates], dtype=float)
    margin = 0.065
    gap = min(0.13, (1 - 2 * margin) / max(1, len(dates) - 1))
    order = np.argsort(desired, kind="stable")
    positions = np.clip(desired[order], margin, 1 - margin)
    for index in range(1, len(positions)):
        positions[index] = max(positions[index], positions[index - 1] + gap)
    positions[-1] = min(positions[-1], 1 - margin)
    for index in range(len(positions) - 2, -1, -1):
        positions[index] = min(positions[index], positions[index + 1] - gap)
    restored = np.empty(len(positions))
    restored[order] = positions
    return [lower + span * float(position) for position in restored]


def _table_cell(sheet, value, kind="body"):
    cell = WriteOnlyCell(sheet, value=value)
    cell.font = Font(name="Calibri", size=10, bold=kind != "body",
                     color="FFFFFF" if kind == "header" else "172B3A")
    cell.alignment = Alignment(vertical="center", horizontal="left")
    if kind == "header":
        cell.fill = PatternFill("solid", fgColor="24476A")
    elif kind == "period":
        cell.fill = PatternFill("solid", fgColor="DFEAF4")
    if isinstance(value, pd.Timestamp):
        cell.value = value.to_pydatetime()
        cell.number_format = "dd-mm-yyyy hh:mm"
    elif isinstance(value, float):
        cell.number_format = "0.0"
    return cell


def _boundary_value(value, is_date=False):
    if value is None or pd.isna(value) or str(value).strip() == "—":
        return "—"
    if is_date:
        parsed = pd.to_datetime(value, format="%d-%m-%Y %H:%M", errors="coerce")
        return parsed if pd.notna(parsed) else str(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return str(value)


def _boundary_tables(sheet, periods, boundaries):
    cells = {}
    bottom = TABLE_START_ROW

    def put(row, column, value, kind="body"):
        cells.setdefault(row, {})[column] = (value, kind)

    def merge(row, first, last):
        sheet.merged_cells.add(f"{get_column_letter(first)}{row}:{get_column_letter(last)}{row}")

    for index, period in enumerate(periods.itertuples(index=False)):
        number = int(period.period)
        top = TABLE_START_ROW + (index // 2) * 15
        first = 1 + (index % 2) * 12
        merge(top, first, first + 10)
        put(top, first, f"Период {number}", "period")
        sheet.row_dimensions[top].height = 24
        for offset, tag, title in ((0, "нач", "Начало: номинальная точка"),
                                   (6, "кон", "Конец: номинальная точка")):
            column = first + offset
            merge(top + 1, column, column + 4)
            put(top + 1, column, title, "header")
            sheet.row_dimensions[top + 1].height = 23
            record = _boundary_record(boundaries, number, tag)
            parameters = [key for key in boundaries.columns if key != "Период"] if boundaries is not None else ["дата"]
            for row_offset, key in enumerate(parameters, start=2):
                row = top + row_offset
                merge(row, column, column + 1)
                merge(row, column + 2, column + 4)
                label = "Дата точки" if key == "дата" else (
                    f"{key}, Па" if key in {"ВЛО", "ФГО", "ФТО", "ФГО+ФТО"}
                    else f"{key}, МВт" if key in {"N", "Nприв"} else key
                )
                put(row, column, label)
                put(row, column + 2, _boundary_value(
                    record[key] if record is not None else None, is_date=key == "дата"
                ))
                sheet.row_dimensions[row].height = 19
                bottom = max(bottom, row)
    return cells, bottom + 2


def write_pressure_sheet(book, headers, header_row, row_count, source: PressureChartData | None):
    """Создаёт один лист с нативной XY-диаграммой и связанными диапазонами."""
    sheet = book.create_sheet("Графики")
    sheet.sheet_view.showGridLines = False
    if source is None:
        return sheet
    frame = source.values.copy().sort_values("date").reset_index(drop=True)
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame = frame.loc[frame["date"].notna()].reset_index(drop=True)
    columns = _source_columns(headers)
    keys = [key for key in pressure_channels(frame) if key in columns]
    if frame.empty or not keys:
        plot = prepare_degradation_plot(
            source.degradation_values, source.periods,
            source.degradation_deltas,
        )
        if plot is None:
            sheet.append(["Нет исходных данных для диаграмм."])
            return sheet
        cells, visible_bottom = add_degradation_chart(sheet, plot, 2)
        sheet.column_dimensions.group("BJ", "BN", hidden=True)
        for column in range(1, 24):
            sheet.column_dimensions[get_column_letter(column)].width = 9.5
        sheet.print_area = f"A1:W{visible_bottom}"
        sheet.append(["Перепады давления: нет исходных данных для диаграммы."])
        for row in range(2, max(visible_bottom, len(plot.points) + 1) + 1):
            values = {
                column: _table_cell(sheet, value, kind)
                for column, (value, kind) in cells.get(row, {}).items()
            }
            if row <= len(plot.points) + 1:
                point = plot.points.iloc[row - 2]
                values[DEGRADATION_HELPER_COLUMN] = point["date"].to_pydatetime()
                for index, key in enumerate(
                    ("power", "fuel_specific", "eta_gtu", "vna"), start=1,
                ):
                    value = point[key]
                    values[DEGRADATION_HELPER_COLUMN + index] = (
                        float(value) if pd.notna(value) else None
                    )
            if row > visible_bottom:
                sheet.row_dimensions[row].hidden = True
            sheet.append([values.get(column) for column in range(1, max(values) + 1)] if values else [])
        return sheet
    series_keys = keys + (["power_n"] if "power_n" in columns and np.isfinite(_number(frame, "power_n")).any() else [])
    pressure_max = max(500, math.ceil(_scale_max(pd.concat(
        [_number(frame, key) for key in keys], ignore_index=True,
    )) / 50) * 50)
    chart = ScatterChart()
    chart.title = None
    chart.roundedCorners = False
    chart.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(noFill=True),
    )
    chart.scatterStyle = "lineMarker"
    chart.width = CHART_WIDTH_CM
    chart.height = CHART_HEIGHT_CM
    chart.x_axis.axId = 10
    chart.y_axis.axId = 20
    chart.x_axis.crossAx = 20
    chart.y_axis.crossAx = 10
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    _axis_title(chart.x_axis, "Дата")
    chart.x_axis.axPos = "b"
    _axis_title(chart.y_axis, "Перепады давления на СФ КВОУ ГТУ, Па")
    chart.x_axis.numFmt = "dd.mm.yyyy"
    chart.x_axis.txPr = _text()
    chart.y_axis.txPr = _text()
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = pressure_max
    chart.y_axis.numFmt = "0"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.majorUnit = math.ceil(pressure_max / 500) * 50
    chart.y_axis.minorUnit = chart.y_axis.majorUnit / 5
    chart.y_axis.majorGridlines = _grid("000000", 12700)
    chart.y_axis.minorGridlines = _grid("A6A6A6", 6350)
    chart.x_axis.majorGridlines = _grid("000000", 12700)
    chart.x_axis.minorGridlines = _grid("A6A6A6", 6350)
    lower, upper = frame["date"].iloc[[0, -1]]
    padding = pd.Timedelta(days=max(1, math.ceil((upper - lower).total_seconds() / 86400 * 0.1)))
    lower, upper = lower.normalize() - padding, upper.normalize() + padding
    chart.x_axis.scaling.min = to_excel(lower.to_pydatetime())
    chart.x_axis.scaling.max = to_excel(upper.to_pydatetime())
    span_days = max(1, (upper - lower).total_seconds() / 86400)
    chart.x_axis.majorUnit = next((days for days in (1, 2, 5, 10, 20, 25, 50, 100, 150, 200, 365) if days >= span_days / 6), 365)
    chart.x_axis.minorUnit = chart.x_axis.majorUnit / 5
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.crosses = "min"
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.legend.txPr = _text()
    chart.legend.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.15, y=0.93, w=0.7, h=0.05,
    ))
    chart.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.075, y=0.035, w=0.845, h=0.80,
    ))
    chart.visible_cells_only = False
    chart.display_blanks = "gap"

    for index, key in enumerate(series_keys, start=1):
        if key == "power_n":
            continue
        title, color = PRESSURE_CHANNELS[key]
        chart.series.append(_point_series(
            sheet, key, f"Перепад давления на {title.removesuffix(', Па')}",
            color, frame, index,
        ))
    helper_points = []
    hide_legend = []
    periods = source.periods if source.periods is not None else pd.DataFrame()
    callout_dates = _callout_dates(periods, lower, upper)

    def add_helper(points, title, **kwargs):
        first = len(helper_points) + 2
        helper_points.extend(points)
        hide_legend.append(len(chart.series))
        chart.series.append(_period_series(sheet, first, points, title, **kwargs))

    for index, period in enumerate(periods.itertuples(index=False)):
        number = int(period.period)
        for moment in (pd.Timestamp(period.date_start), pd.Timestamp(period.date_end)):
            points = [(moment, 0.0), (moment, pressure_max)]
            add_helper(points, f"Граница периода {number}")
        arrow_y = pressure_max * 0.94
        add_helper([(pd.Timestamp(period.date_start), arrow_y),
                    (pd.Timestamp(period.date_end), arrow_y)],
                   f"Диапазон периода {number}", arrow=True)
        midpoint = pd.Timestamp(period.date_start) + (pd.Timestamp(period.date_end) - pd.Timestamp(period.date_start)) / 2
        add_helper([(midpoint, arrow_y)],
                   f"{number} период исследования\n"
                   f"{period.date_start:%d.%m.%Y} - {period.date_end:%d.%m.%Y}", label=True)
        for offset, tag in enumerate(("нач", "кон")):
            add_helper([(callout_dates[index * 2 + offset], pressure_max * 0.78)],
                       _callout(_boundary_record(source.boundaries, number, tag), keys), label=True)
    grouped_events = {}
    for event in source.events:
        moment = pd.Timestamp(event.date)
        if lower <= moment <= upper:
            grouped_events.setdefault(moment, []).append(event.label)
    for moment, labels in sorted(grouped_events.items()):
        title = " / ".join(labels)
        add_helper([(moment, 0.0), (moment, pressure_max)], f"Событие: {title}")
        add_helper([(moment, pressure_max * 0.52)], title, label=True, vertical=True)
    chart.legend.legendEntry = [LegendEntry(idx=index, delete=True) for index in hide_legend]

    if "power_n" in series_keys:
        power = ScatterChart()
        power.scatterStyle = "lineMarker"
        power.x_axis = NumericAxis(axId=40, crossAx=30, axPos="b", delete=True)
        power.x_axis.majorGridlines = None
        power.x_axis.scaling.min = chart.x_axis.scaling.min
        power.x_axis.scaling.max = chart.x_axis.scaling.max
        power.y_axis.axId = 30
        power.y_axis.crossAx = 40
        power.y_axis.delete = False
        power.y_axis.axPos = "r"
        _axis_title(power.y_axis, "Мощность ГТУ, МВт")
        power.y_axis.crosses = "max"
        max_power = _number(frame, "power_n").max()
        power.y_axis.scaling.min = 0 if max_power < 50 else 50
        power.y_axis.scaling.max = 100 if max_power <= 100 else _scale_max(_number(frame, "power_n"))
        power.y_axis.majorUnit = 5 if power.y_axis.scaling.max <= 100 else 10
        power.y_axis.numFmt = "0"
        power.y_axis.tickLblPos = "nextTo"
        power.y_axis.txPr = _text()
        power.y_axis.majorGridlines = None
        power.series.append(_point_series(sheet, "power_n", "Nакт", POWER_COLOR,
                                          frame, series_keys.index("power_n") + 1))
        chart += power
    sheet.add_chart(chart, "A2")
    for column in range(1, 24):
        sheet.column_dimensions[get_column_letter(column)].width = 9.5
    sheet.column_dimensions.group("AX", "BH", hidden=True)
    table_cells, visible_bottom = _boundary_tables(sheet, periods, source.boundaries)
    degradation_plot = prepare_degradation_plot(
        source.degradation_values, periods, source.degradation_deltas,
    )
    if degradation_plot is not None:
        degradation_cells, degradation_bottom = add_degradation_chart(
            sheet, degradation_plot, start_row=1,
            start_col=DEGRADATION_CHART_COLUMN,
        )
        for row, cells in degradation_cells.items():
            table_cells.setdefault(row, {}).update(cells)
        visible_bottom = max(visible_bottom, degradation_bottom)
        sheet.column_dimensions.group("BJ", "BN", hidden=True)
        for column in range(DEGRADATION_CHART_COLUMN, VISIBLE_LAST_COLUMN + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 9.5
        sheet.column_dimensions[get_column_letter(VISIBLE_LAST_COLUMN)].width = 12
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 2 if degradation_plot is not None else 1
    sheet.page_setup.fitToHeight = 0
    last_visible_column = (
        get_column_letter(VISIBLE_LAST_COLUMN)
        if degradation_plot is not None else "W"
    )
    sheet.print_area = f"A1:{last_visible_column}{visible_bottom}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    last_row = max(
        visible_bottom, len(frame) + 1, len(helper_points) + 1,
        len(degradation_plot.points) + 1 if degradation_plot is not None else 0,
    )
    for row in range(1, last_row + 1):
        values = {}
        if row in table_cells:
            values.update({column: _table_cell(sheet, value, kind)
                           for column, (value, kind) in table_cells[row].items()})
        if row == 1:
            for index, title in enumerate(["Дата"] + [PRESSURE_CHANNELS[key][0] if key in PRESSURE_CHANNELS else "N, МВт" for key in series_keys]):
                values[HELPER_START_COLUMN + index] = title
            values[PERIOD_HELPER_COLUMN] = "Даты границ и подписей"
            values[PERIOD_HELPER_COLUMN + 1] = "Координата по оси Па"
            if degradation_plot is not None:
                for index, title in enumerate(DEGRADATION_HELPER_HEADERS):
                    values[DEGRADATION_HELPER_COLUMN + index] = title
        if 2 <= row <= len(frame) + 1:
            record = frame.iloc[row - 2]
            excel_row = record.get("excel_row")
            linked = pd.notna(excel_row) and "date" in columns
            if linked:
                excel_row = int(excel_row)
                if not header_row + 2 <= excel_row <= row_count:
                    raise ValueError("Строка диаграммы вышла за диапазон листа «Данные».")
                date_ref = f"'Данные'!{get_column_letter(columns['date'])}{excel_row}"
                values[HELPER_START_COLUMN] = f"={date_ref}"
            else:
                values[HELPER_START_COLUMN] = record["date"].to_pydatetime()
            for index, key in enumerate(series_keys, start=1):
                if linked:
                    ref = f"'Данные'!{get_column_letter(columns[key])}{excel_row}"
                    values[HELPER_START_COLUMN + index] = f"=IF(ISNUMBER({ref}),IF({ref}>=0,{ref},NA()),NA())"
                else:
                    value = pd.to_numeric(record.get(key), errors="coerce")
                    values[HELPER_START_COLUMN + index] = float(value) if pd.notna(value) and np.isfinite(value) and value >= 0 else "=NA()"
        if 2 <= row <= len(helper_points) + 1:
            moment, value = helper_points[row - 2]
            values[PERIOD_HELPER_COLUMN] = moment.to_pydatetime()
            values[PERIOD_HELPER_COLUMN + 1] = value
        if degradation_plot is not None and 2 <= row <= len(degradation_plot.points) + 1:
            point = degradation_plot.points.iloc[row - 2]
            values[DEGRADATION_HELPER_COLUMN] = point["date"].to_pydatetime()
            for index, key in enumerate(("power", "fuel_specific", "eta_gtu", "vna"), start=1):
                value = point[key]
                values[DEGRADATION_HELPER_COLUMN + index] = (
                    float(value) if pd.notna(value) else None
                )
        if row > visible_bottom:
            sheet.row_dimensions[row].hidden = True
        if values:
            sheet.append([values.get(column) for column in range(1, max(values) + 1)])
        else:
            sheet.append([])
    return sheet
