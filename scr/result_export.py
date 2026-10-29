"""Книга результатов: исходные значения, таблицы расчёта и графики."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import os
import tempfile
from contextlib import suppress

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from data_assembly import AssembledData
from data_loader import _find_header_row
from pressure_chart import PressureChartData, write_pressure_sheet


@dataclass
class DataSheetValues:
    """Снимок первого листа книги, включая исходную шапку и порядок строк."""

    values: pd.DataFrame
    header_row: int

    def apply_correction_profile(self, profile):
        """Применяет все именованные коэффициенты; готовую термодинамику сохраняет."""
        from data_assembly import calculate_data_columns
        from openpyxl.utils import column_index_from_string
        frame = self.values.iloc[self.header_row + 1:].copy()
        frame.columns = [get_column_letter(i+1) for i in range(len(frame.columns))]
        selected = frame.loc[pd.to_datetime(frame["B"],errors="coerce").notna()].copy()
        # Та же поддержка старого расположения, что и в assemble_sensor_data,
        # но только в рабочей копии: исходная AC и AD в экспорте не меняются.
        if "AD" in selected and pd.isna(self.values.iat[self.header_row,29]):
            shifted = selected["AC"].isna() & selected["AD"].notna()
            if shifted.any():
                selected["AC"] = selected["AC"].astype(object)
                selected.loc[shifted,"AC"] = selected.loc[shifted,"AD"]
                print(f"Для поправки давления ух.газов использована AD вместо пустой AC: {int(shifted.sum())} строк; исходные колонки сохранены.")
        data = AssembledData(selected, {}, 0, 0, {}, [])
        calculate_data_columns(data, profile, thermodynamics=False)
        outputs = ("D","H","N","O","S","T","U","Y","AE","AI","AN",
                   "AJ","AK","AL","AQ","AT","AV","AX","AR","AU","AW")
        for column in outputs:
            index = column_index_from_string(column)-1
            if index >= len(self.values.columns):
                raise ValueError(f"В принятой структуре СызТЭЦ нет колонки {column}.")
            destination = self.values.columns[index]
            self.values[destination] = self.values[destination].astype(object)
            self.values.loc[selected.index,destination] = selected[column].to_numpy()

    def apply_gas_profile(self, profile):
        """Обновляет газозависимые столбцы только датированных строк текущей книги."""
        from data_assembly import recalculate_gas_columns
        frame = self.values.iloc[self.header_row + 1:].copy()
        frame.columns = [get_column_letter(i + 1) for i in range(len(frame.columns))]
        dated = pd.to_datetime(frame["B"], errors="coerce").notna()
        selected = frame.loc[dated].copy()
        recalculate_gas_columns(selected, profile)
        outputs = ("AJ", "AK", "AL", "AQ", "AT", "AV", "AX", "AR", "AU", "AW")
        from openpyxl.utils import column_index_from_string
        for column in outputs:
            index = column_index_from_string(column) - 1
            if index >= len(self.values.columns):
                raise ValueError(f"В структуре книги СызТЭЦ нет расчётной колонки {column}.")
            destination = self.values.columns[index]
            self.values[destination] = self.values[destination].astype(object)
            self.values.loc[selected.index, destination] = selected[column].to_numpy()

    def as_headered_frame(self) -> pd.DataFrame:
        frame = self.values.iloc[self.header_row + 1:].copy()
        frame.columns = self.values.iloc[self.header_row].tolist()
        frame = frame.reset_index(drop=True).infer_objects()
        frame.attrs["excel_header_row"] = self.header_row
        return frame


def read_data_sheet(path: str | Path) -> DataSheetValues:
    with pd.ExcelFile(path, engine="calamine") as book:
        if not book.sheet_names or book.sheet_names[0] != "Данные":
            raise ValueError("Первый лист основной книги должен называться «Данные».")
        values = book.parse("Данные", header=None)
    header_row = _find_header_row(values.head(25))
    if header_row is None:
        raise ValueError("На листе «Данные» не найдена строка заголовков.")
    return DataSheetValues(values, header_row)


def collect_result_tables(
    periods: pd.DataFrame,
    output_tables: dict,
    vna_results: list[pd.DataFrame],
    economics_table: pd.DataFrame,
) -> list[tuple[str, pd.DataFrame | None]]:
    """Все строки табличных результатов без экранных разделителей колонок."""
    sections: list[tuple[str, pd.DataFrame | None]] = []
    if periods is not None and not periods.empty:
        period_columns = {
            "period": "Период", "date_start": "Начало", "date_end": "Конец",
            "records": "Количество записей", "duration_days": "Длительность, суток",
        }
        selected = [col for col in period_columns if col in periods]
        sections.append(("Периоды расчёта", periods[selected].rename(columns=period_columns)))
    sections.append(("Параметры на границах периодов", output_tables.get("boundary")))
    for index, table in enumerate(vna_results, start=1):
        period = table.attrs.get("period_number", index)
        sections.append((
            f"Период {period}: TOP-10 интервалов ВНА",
            table.drop(columns=["share_percent"], errors="ignore"),
        ))
    if not vna_results:
        sections.append(("TOP-10 интервалов ВНА", None))
    sections.extend([
        ("Дельты по регрессиям", output_tables.get("delta")),
        ("Наработка и пороги", output_tables.get("nominal_hours")),
        ("Экономика", economics_table),
    ])
    return sections


_HEADER_FILL = PatternFill("solid", fgColor="24476A")
_TITLE_FILL = PatternFill("solid", fgColor="DFEAF4")
_ALT_FILL = PatternFill("solid", fgColor="F4F7FA")
_VALUE_FONT = Font(name="Calibri", size=11, color="172B3A")
_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="172B3A")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_VALUE_ALIGNMENT = Alignment(vertical="center")
_WRAP_ALIGNMENT = Alignment(vertical="center", wrap_text=True)


def _cell(sheet, value, *, header=False, title=False, total=False, alternate=False):
    if isinstance(value, np.generic):
        value = value.item()
    if value is None or pd.isna(value):
        value = None
    elif isinstance(value, float) and not np.isfinite(value):
        value = None
    elif isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    elif isinstance(value, pd.Timedelta):
        value = value.to_pytimedelta()
    cell = WriteOnlyCell(sheet, value=value)
    # Значения, начинающиеся с '=', остаются текстом, а не формулой Excel.
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    cell.font = _HEADER_FONT if header else _BOLD_FONT if title or total else _VALUE_FONT
    cell.alignment = _WRAP_ALIGNMENT if header or title else _VALUE_ALIGNMENT
    if header:
        cell.fill = _HEADER_FILL
    elif title or total:
        cell.fill = _TITLE_FILL
    elif alternate:
        cell.fill = _ALT_FILL
    if isinstance(value, (datetime, date)):
        cell.number_format = "dd-mm-yyyy hh:mm"
    elif isinstance(value, float):
        cell.number_format = "0.000"
    elif isinstance(value, int):
        cell.number_format = "0"
    return cell


def _write_data_sheet(book: Workbook, data: AssembledData | DataSheetValues) -> tuple:
    sheet = book.create_sheet("Данные")
    if isinstance(data, DataSheetValues):
        frame = data.values
        header_row = data.header_row
        rows = frame.itertuples(index=False, name=None)
        headers = frame.iloc[header_row].tolist()
        row_count, column_count = frame.shape
    else:
        from openpyxl.utils import column_index_from_string
        column_count = max(column_index_from_string(col) for col in data.values)
        letters = [get_column_letter(index) for index in range(1, column_count + 1)]
        headers = [data.headers.get(col) for col in letters]
        header_row = 0
        from itertools import chain
        extra = data.undated_rows
        rows = chain(
            [headers], data.values.reindex(columns=letters).itertuples(index=False, name=None),
            extra.reindex(columns=letters).itertuples(index=False, name=None)
            if extra is not None else (),
        )
        row_count = len(data.values) + 1 + (len(extra) if extra is not None else 0)
    sheet.freeze_panes = f"B{header_row + 2}"
    sheet.auto_filter.ref = f"A{header_row + 1}:{get_column_letter(column_count)}{row_count}"
    sheet.sheet_view.showGridLines = False
    for index, header in enumerate(headers, start=1):
        text = "" if pd.isna(header) else str(header)
        sheet.column_dimensions[get_column_letter(index)].width = (
            22 if index == 2 else min(26, max(13, len(text) + 2))
        )
    sheet.row_dimensions[header_row + 1].height = 42
    for index, row in enumerate(rows):
        sheet.append([_cell(sheet, value, header=index == header_row) for value in row])
    return headers, header_row, row_count


def _write_tables_sheet(book: Workbook, sections: list[tuple[str, pd.DataFrame | None]]) -> None:
    sheet = book.create_sheet("Таблицы")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A3"
    max_columns = max((len(table.columns) for _, table in sections if table is not None), default=1)
    for index in range(max_columns):
        width = 14
        for _, table in sections:
            if table is None or index >= len(table.columns):
                continue
            values = table.iloc[:, index]
            width = max(width, min(28, len(str(table.columns[index])) + 2))
            if pd.api.types.is_datetime64_any_dtype(values):
                width = max(width, 22)
            elif not pd.api.types.is_numeric_dtype(values):
                width = max(width, min(28, max((len(str(v)) + 2 for v in values.dropna()), default=0)))
        sheet.column_dimensions[get_column_letter(index + 1)].width = width
    row_number = 0
    for title, table in sections:
        count = max(1, len(table.columns)) if table is not None else 1
        sheet.merged_cells.add(
            f"A{row_number + 1}:{get_column_letter(max(count, 3))}{row_number + 1}"
        )
        sheet.row_dimensions[row_number + 1].height = 30
        sheet.append([_cell(sheet, title, title=True)] + [
            _cell(sheet, None, title=True) for _ in range(max(count, 3) - 1)
        ])
        row_number += 1
        if table is None or table.empty:
            sheet.append([_cell(sheet, "Нет данных")])
            row_number += 1
        else:
            sheet.row_dimensions[row_number + 1].height = 60
            sheet.append([_cell(sheet, str(col), header=True) for col in table.columns])
            row_number += 1
            for index, row in enumerate(table.itertuples(index=False, name=None)):
                total = any(
                    isinstance(value, str) and (
                        value.strip().upper() == "ИТОГО" or value.startswith("Средневзвеш")
                    ) for value in row
                )
                sheet.append([_cell(sheet, value, total=total, alternate=index % 2 == 1) for value in row])
                row_number += 1
        sheet.append([])
        sheet.append([])
        row_number += 2


def save_result_workbook(
    data: AssembledData | DataSheetValues,
    sections: list[tuple[str, pd.DataFrame | None]],
    output_path: str | Path,
    *, pressure_chart: PressureChartData | None = None,
    correction_sheet=None, gas_passport=None, correction_profile=None,
) -> Path:
    """Пишет новую книгу атомарно, сохраняя прежний файл при ошибке экспорта."""
    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Книга результатов должна иметь расширение .xlsx.")
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    temporary = None
    try:
        layout = _write_data_sheet(book, data)
        _write_tables_sheet(book, sections)
        write_pressure_sheet(book, *layout, pressure_chart)
        from correction_sheet import write_correction_sheet
        write_correction_sheet(book, correction_sheet, gas_passport, correction_profile)
        with tempfile.NamedTemporaryFile(prefix=".gtu_result_", suffix=".xlsx", dir=path.parent, delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        for sheet in book.worksheets:
            if not sheet.closed:
                with suppress(Exception):
                    sheet.close()
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return path
