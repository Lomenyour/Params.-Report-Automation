"""Одна нативная расходная характеристика на период из существующих кандидатов."""
from dataclasses import dataclass
import math

import numpy as np
import pandas as pd
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, NumData, NumDataSource, NumVal
from openpyxl.chart.series import SeriesLabel, XYSeries
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineEndProperties, LineProperties
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill

from chart_text import rich_text, label_series, fitted_label_layout
from chart_layout import cell_anchor, cell_size_points
from RH_raschet import select_rh_variants, rh_equation


@dataclass
class RHPlot:
    result: dict
    start: pd.DataFrame
    end: pd.DataFrame


def _arrow(start, end, color):
    series = XYSeries(tx=SeriesLabel(v="Стрелка состояния"),
        xVal=AxDataSource(numLit=NumData(ptCount=2, pt=[NumVal(idx=i, v=float(point[0])) for i, point in enumerate((start,end))])),
        yVal=NumDataSource(numLit=NumData(ptCount=2, pt=[NumVal(idx=i, v=float(point[1])) for i, point in enumerate((start,end))])))
    series.marker.symbol = "none"
    series.graphicalProperties.line = LineProperties(solidFill=color, w=19050,
        tailEnd=LineEndProperties(type="triangle", w="sm", len="sm"))
    return series


def prepare_rh_plots(values, results):
    if values is None or values.empty or not results:
        return []
    plots = []
    for result in select_rh_variants(results):
        parts = []
        for prefix in ("start", "end"):
            dates = pd.to_datetime(values["date"], errors="coerce")
            part = values.loc[(values["period"] == result["period"])
                & (dates >= result[f"{prefix}_window_begin"])
                & (dates <= result[f"{prefix}_window_end"])].copy()
            part = part.sort_values("date", kind="stable")
            if part.empty:
                break
            for key in ("power", "fuel"):
                part[key] = pd.to_numeric(part[key], errors="coerce")
                if not np.isfinite(part[key]).all():
                    raise ValueError(f"Период {result['period']}: недопустимые данные РХ в {prefix}.")
            parts.append(part)
        if len(parts) == 2:
            plots.append(RHPlot(result, *parts))
    return plots


def _bounds(values):
    low, high = float(np.min(values)), float(np.max(values))
    span = max(high - low, max(abs(low), abs(high)) * .01, .1)
    magnitude = 10 ** math.floor(math.log10(span / 6))
    step = next(unit * magnitude for unit in (1, 2, 5, 10) if unit * magnitude >= span / 6)
    return math.floor((low - .1 * span) / step) * step, math.ceil((high + .1 * span) / step) * step, step


def _series(sheet, x_column, y_column, first, last, x_values, y_values, title, color, line=False):
    series = Series(Reference(sheet, min_col=y_column, min_row=first, max_row=last),
                    Reference(sheet, min_col=x_column, min_row=first, max_row=last), title=title)
    for cache, values in ((series.xVal.numRef, x_values), (series.yVal.numRef, y_values)):
        cache.numCache = NumData(ptCount=len(values), pt=[NumVal(idx=i, v=float(v)) for i, v in enumerate(values)])
    if line:
        series.graphicalProperties.line.solidFill = color
        series.marker.symbol = "none"
        series.graphicalProperties.line.width = 19050
    else:
        series.graphicalProperties.line.noFill = True
        series.marker.symbol = "circle"
        series.marker.size = 3
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
    return series


def _write_data_blocks(book, plots):
    sheet = book.create_sheet("Данные РХ")
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    for column in ("A", "B", "C", "D", "E", "G", "H", "I"):
        sheet.column_dimensions[column].width = 25
    blocks, offset = {}, 0

    def cell(value, *, header=False):
        item = WriteOnlyCell(sheet, value=value)
        item.font = Font(name="Calibri", size=11, bold=header,
                         color="FFFFFF" if header else "000000")
        item.alignment = Alignment(vertical="center", wrap_text=True)
        if header:
            item.fill = PatternFill("solid", fgColor="24476A")
        if isinstance(value, pd.Timestamp):
            item.value = value.to_pydatetime()
            item.number_format = "dd.mm.yyyy hh:mm"
        elif isinstance(value, float):
            item.number_format = "0.000"
        return item

    for plot in plots:
        result = plot.result
        number = int(result["period"])
        first = offset + 3
        sheet.merged_cells.add(f"A{offset + 1}:I{offset + 1}")
        sheet.row_dimensions[offset + 1].height = 24
        sheet.row_dimensions[offset + 2].height = 36
        sheet.append([cell(f"Период {number}: РХ, вариант {result['rank']}", header=True)])
        sheet.append([cell(value, header=True) for value in
            ("Состояние", "Дата", "Nприв, МВт", "Bприв, нм³/ч", "Строка исходного Excel",
             None, "N линии", "Bчист", "Bзагр")])
        combined = pd.concat([plot.start, plot.end], ignore_index=True)
        x = np.linspace(combined.power.min(), combined.power.max(), 100)
        clean = result["k_start"] * x + result["b_start"]
        dirty = result["k_end"] * x + result["b_end"]
        length = max(len(combined), len(x))
        for i in range(length):
            row = [None] * 9
            if i < len(combined):
                point = combined.iloc[i]
                row[:5] = ["Условно-чистое" if i < len(plot.start) else "Условно-загрязнённое",
                    pd.Timestamp(point["date"]), float(point.power), float(point.fuel),
                    int(point.excel_row) if "excel_row" in point and pd.notna(point.excel_row) else None]
            if i < len(x):
                row[6:] = [float(x[i]), float(clean[i]), float(dirty[i])]
            sheet.append([cell(value) for value in row])
        blocks[number] = (first, combined, x, clean, dirty)
        sheet.append([])
        sheet.append([])
        offset += length + 4
    return sheet, blocks


def add_rh_charts(sheet, plots, start_row, *, two_columns=True):
    cells = {}
    bottom = start_row
    if not plots:
        return cells, bottom
    data_sheet, blocks = _write_data_blocks(sheet.parent, plots)
    for index, plot in enumerate(plots):
        result = plot.result
        number = int(result["period"])
        first_data, combined, x, clean, dirty = blocks[number]

        column = 1 + (index % 2) * 16 if two_columns else 1
        top = start_row + (index // 2 if two_columns else index) * 38
        cells.setdefault(top, {})[column] = (f"Период {number}: расходная характеристика, вариант {result['rank']}", "period")
        sheet.merged_cells.add(f"{get_column_letter(column)}{top}:{get_column_letter(column + 13)}{top}")
        sheet.row_dimensions[top].height = 24
        chart = ScatterChart()
        chart.title = None
        chart.scatterStyle = "lineMarker"
        chart.graphical_properties = GraphicalProperties(solidFill="FFFFFF", ln=LineProperties(noFill=True))
        for axis, title in ((chart.x_axis, "Приведённая мощность, МВт"),
                            (chart.y_axis, "Приведённый расход газа в КС ГТУ, нм³/ч")):
            axis.title = title
            axis.title.tx.rich = rich_text(title, bold=True)
            axis.title.txPr = rich_text("", bold=True)
            axis.txPr = rich_text("")
            axis.delete = False
            axis.numFmt = "0" if axis is chart.y_axis else "0.0"
            axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="808080", w=9525)))
            axis.minorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="C0C0C0", w=4762)))
        chart.x_axis.scaling.min, chart.x_axis.scaling.max, chart.x_axis.majorUnit = _bounds(x)
        chart.y_axis.scaling.min, chart.y_axis.scaling.max, chart.y_axis.majorUnit = _bounds(np.r_[combined.fuel, clean, dirty])
        chart.x_axis.minorUnit = chart.x_axis.majorUnit / 5
        chart.y_axis.minorUnit = chart.y_axis.majorUnit / 5
        chart.x_axis.tickLblPos = "low"
        chart.legend.position = "b"
        chart.legend.overlay = False
        chart.legend.txPr = rich_text("")
        chart.legend.layout = Layout(manualLayout=ManualLayout(xMode="edge", yMode="edge", x=.1, y=.94, w=.8, h=.05))
        chart.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                       x=.16, y=.04, w=.79, h=.79))
        for part, first, title, color in ((plot.start, first_data, "Условно-чистая", "4472C4"),
                (plot.end, len(plot.start) + first_data, "Условно-загрязнённая", "C00000")):
            chart.series.append(_series(data_sheet, 3, 4, first, first + len(part) - 1,
                                        part.power, part.fuel, title, color))
        chart.series.append(_series(data_sheet, 7, 8, first_data, first_data + len(x) - 1,
                                    x, clean, "Линия чистая", "4472C4", True))
        chart.series.append(_series(data_sheet, 7, 9, first_data, first_data + len(x) - 1,
                                    x, dirty, "Линия загрязнённая", "C00000", True))
        xlow, xhigh = chart.x_axis.scaling.min, chart.x_axis.scaling.max
        ylow, yhigh = chart.y_axis.scaling.min, chart.y_axis.scaling.max
        for prefix, name, state, xr, yr, target_fraction, color in (
            ("end", "Bзагр", "Условно-загрязнённое", .36, .84, .50, "C00000"),
            ("start", "Bчист", "Условно-чистое", .65, .21, .88, "4472C4"),
        ):
            text = (f'"{state}"\nсостояние {result[f"{prefix}_window_begin"]:%d.%m.%Y}\n'
                    f"{rh_equation(name, result[f'k_{prefix}'], result[f'b_{prefix}'])}")
            target_x = float(x[0] + (x[-1] - x[0]) * target_fraction)
            target_y = result[f"k_{prefix}"] * target_x + result[f"b_{prefix}"]
            chart.series.append(_arrow((xlow + (xhigh-xlow)*xr, ylow + (yhigh-ylow)*yr),
                                       (target_x, target_y), color))
            chart.series.append(label_series(xlow + (xhigh - xlow) * xr,
                ylow + (yhigh - ylow) * yr, text, bold=True, color=color, size=1400,
                layout=fitted_label_layout(text, cell_size_points(sheet,column,top+2,14,33),
                                           (.16+.79*xr, .04+.79*(1-yr)))))
        chart.legend.legendEntry = [LegendEntry(idx=i, delete=True) for i in range(2, len(chart.series))]
        sheet.add_chart(chart, cell_anchor(column, top + 2, 14, 33))
        bottom = max(bottom, top + 35)
    return cells, bottom
