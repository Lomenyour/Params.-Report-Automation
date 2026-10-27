"""Порог только граничной таблицы, без изменения расчётных периодов."""
from pathlib import Path
import io
import sys
import unittest
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scr"))
from RH_raschet import build_boundary_table, boundary_vna_rows
from pressure_chart import pressure_display_periods, PressureChartData, write_pressure_sheet


class BoundaryCoverageTest(unittest.TestCase):
    def setUp(self):
        self.values = pd.DataFrame([
            {"period": period, "date": pd.Timestamp(2026, period, 1) + pd.Timedelta(minutes=30 * index),
             "power": 100., "power_n": 80., "vna": vna, "vlo": float(index)}
            for period, values in ((1, [20., 82.607, 83., 84., 86., 86.5, 87., 88., 89., 30.]),
                                   (2, [10., 70., 71., 72., 73., 74., 75., 76., 77., 20.]))
            for index, vna in enumerate(values)
        ])
        self.periods = pd.DataFrame([{"period": p, "date_start": g.date.min(), "date_end": g.date.max()}
                                     for p, g in self.values.groupby("period")])

    def test_each_period_has_its_own_threshold_and_periods_are_unchanged(self):
        original = self.values.copy(deep=True)
        periods = self.periods.copy(deep=True)
        table = build_boundary_table(self.values, self.periods)
        self.assertEqual(table["Порог ВНА, %"].tolist(), [82.6, 82.6, 70., 70.])
        self.assertEqual(table["Покрытие, %"].tolist(), [80.] * 4)
        self.assertEqual(table["дата"].tolist(), ["01-01-2026 00:30", "01-01-2026 04:00",
                                                  "01-02-2026 00:30", "01-02-2026 04:00"])
        display = pressure_display_periods(self.periods, table)
        self.assertEqual(display.iloc[0].date_start, pd.Timestamp("2026-01-01 00:30"))
        self.assertEqual(display.iloc[0].date_end, pd.Timestamp("2026-01-01 04:00"))
        pd.testing.assert_frame_equal(self.values, original)
        pd.testing.assert_frame_equal(self.periods, periods)

    def test_invalid_vna_is_reported_and_never_replaced_by_zero(self):
        self.values.loc[[0, 1, 2], "vna"] = [np.nan, -1., 101.]
        table = build_boundary_table(self.values, self.periods)
        summary = table.attrs["vna_boundary_summary"][1]
        self.assertEqual(summary["valid_records"], 7)
        self.assertEqual(summary["invalid_records"], 3)
        self.assertEqual(summary["availability_percent"], 70.)
        self.assertGreaterEqual(summary["coverage_percent"], 80.)
        self.assertTrue(pd.isna(self.values.loc[0, "vna"]))

    def test_stopped_unassigned_or_other_period_rows_do_not_affect_threshold(self):
        self.values.loc[0, ["power", "vna"]] = [0., 100.]
        self.values.loc[1, ["period", "vna"]] = [0, 100.]
        row = self.periods.iloc[0]
        selected, summary = boundary_vna_rows(self.values, row.date_start, row.date_end, 1)
        self.assertEqual(summary["working_records"], 8)
        self.assertTrue((selected.period == 1).all())
        self.assertTrue((selected.power > 10).all())
        self.assertTrue((selected.vna >= summary["threshold"]).all())

    def test_no_valid_vna_keeps_missing_boundary_and_no_fictitious_chart_dates(self):
        self.values.loc[self.values.period == 1, "vna"] = np.nan
        table = build_boundary_table(self.values, self.periods)
        self.assertEqual(table.iloc[0]["дата"], "—")
        self.assertTrue(pd.isna(table.iloc[0]["Порог ВНА, %"]))
        display = pressure_display_periods(self.periods, table)
        self.assertEqual(display.period.tolist(), [2])

    def test_all_missing_vna_still_exports_pressure_chart_without_fictitious_dates(self):
        self.values["vna"] = np.nan
        table = build_boundary_table(self.values, self.periods)
        book = Workbook(write_only=True)
        headers = ["Дата", "ВЛО, Па", "N, МВт"]
        data = book.create_sheet("Данные")
        data.append(headers)
        for row in self.values.itertuples():
            data.append([row.date.to_pydatetime(), row.vlo, row.power_n])
        write_pressure_sheet(book, headers, 0, len(self.values) + 1,
            PressureChartData(self.values, self.periods, table))
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        saved = load_workbook(output)
        chart = saved["Графики"]._charts[0]
        self.assertEqual(len(chart.series), 1)  # Только ВЛО, без ложных линий границ.
        self.assertIsNone(chart.series[0].dLbls)
        self.assertTrue(self.values.vna.isna().all())


if __name__ == "__main__":
    unittest.main()
