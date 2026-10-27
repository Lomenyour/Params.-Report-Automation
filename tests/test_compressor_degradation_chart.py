"""TOP-1 каждого периода определяет собственную выборку для Excel-графика."""
import io
import sys
import unittest
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scr"))
from compressor_degradation_chart import (  # noqa: E402
    CompressorPlot, add_compressor_charts, prepare_compressor_plots,
)
from pressure_chart import PressureChartData, write_pressure_sheet  # noqa: E402


class CompressorDegradationChartTest(unittest.TestCase):
    def test_period_filters_and_exported_charts(self):
        records = []
        results = []
        for period in range(1, 4):
            first = pd.Timestamp(2026, period, 1)
            for index, (vna, t, pressure, eta) in enumerate((
                (86.21, 10.1, 990.1, 84.01),
                (86.29, 10.4, 990.4, 84.02),
                (86.31, 10.2, 990.2, 84.01),  # вне TOP-1 ВНА
                (86.22, 11.0, 990.2, 84.01),  # вне температуры
                (86.23, 10.2, 991.0, 84.01),  # вне давления
                (86.23, 10.2, 990.2, 84.89),  # только КПД вне диапазона
                (86.23, 10.2, 990.2, 82.86),  # только КПД вне диапазона
            )):
                records.append({
                    "period": period, "date": first + pd.Timedelta(hours=index),
                    "vna": vna, "t": t, "p": pressure, "tkk": 370.0,
                    "pkk": 13.0, "pi_k": 13.1 + index / 10,
                    "eta_compressor": eta, "excel_row": (period - 1) * 7 + index + 2,
                    "vlo": 200.0, "fgo_fto": 225.0, "power_n": 85.0,
                })
            table = pd.DataFrame([{
                "vna_interval": "86.2-86.3", "t_1": 10.1, "t_2": 10.4,
                "p_1": 990.1, "p_2": 990.4,
                "date_1": first, "date_2": first + pd.Timedelta(hours=1),
                "eta_1": 84.01, "eta_2": 84.02,
                "pi_k_1": 13.1, "pi_k_2": 13.2,
            }])
            table.attrs["period_number"] = period
            results.append(table)
        values = pd.DataFrame(records)
        periods = pd.DataFrame([{
            "period": period,
            "date_start": pd.Timestamp(2026, period, 1),
            "date_end": pd.Timestamp(2026, period, 1, 6),
        } for period in range(1, 4)])

        plots = prepare_compressor_plots(values, periods, results)
        self.assertEqual([len(plot.points) for plot in plots], [2, 2, 2])
        self.assertEqual(plots[0].points["pi_k"].tolist(), [13.1, 13.2])
        self.assertEqual(plots[0].points["eta"].round(2).tolist(), [84.01, 84.02])

        book = Workbook(write_only=True)
        data_sheet = book.create_sheet("Данные")
        headers = ["Дата", "ВЛО, Па", "ФГО+ФТО, Па", "N, МВт"]
        data_sheet.append(headers)
        for row in values.itertuples(index=False):
            data_sheet.append([row.date.to_pydatetime(), row.vlo,
                               row.fgo_fto, row.power_n])
        write_pressure_sheet(
            book, headers, 0, len(values) + 1,
            PressureChartData(values, periods, None, vna_top10=results,
                              degradation_values=values),
        )
        output = io.BytesIO()
        book.save(output)

        # openpyxl не восстанавливает оформление plotArea при чтении;
        # проверяем именно XML, который получит Excel.
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        with ZipFile(output) as archive:
            for index in range(3, 6):
                root = ET.fromstring(archive.read(f"xl/charts/chart{index}.xml"))
                fill = root.find("c:chart/c:plotArea/c:spPr/a:solidFill/a:srgbClr", ns)
                self.assertEqual(fill.get("val"), "FFFFFF")
                rich_labels = root.findall(".//c:dLbl/c:tx/c:rich", ns)
                self.assertEqual(len(rich_labels), 3)
                for rich in rich_labels:
                    for paragraph in rich.findall("a:p", ns):
                        style = paragraph.find("a:pPr/a:defRPr", ns)
                        self.assertEqual(style.get("sz"), "1000")
                        self.assertEqual(style.find("a:latin", ns).get("typeface"), "Times New Roman")
                self.assertTrue(root.findall(".//a:rPr[@baseline='-25000']", ns))
                criteria_lines = ["".join(node.itertext())
                                  for node in rich_labels[-1].findall("a:p", ns)]
                self.assertEqual(criteria_lines, [
                    "86,2 % ≤ ВНА ≤ 86,3 % = const.",
                    "10,1 °С ≤ tнк ≤ 10,4 °С",
                    "990,1 мбар ≤ Pнк ≤ 990,4 мбар",
                ])
        output.seek(0)
        saved = load_workbook(output)
        sheet = saved["Графики"]
        self.assertEqual(saved.sheetnames[-3:], ["Компрессор П1", "Компрессор П2", "Компрессор П3"])
        self.assertEqual(len(sheet._charts), 5)  # Перепады, общая деградация, 3 периода
        for number, chart in enumerate(sheet._charts[2:], start=1):
            self.assertEqual(len(chart._charts), 2)
            self.assertEqual(len(chart.series), 6)  # КПД и пять элементов подписей
            source = saved[f"Компрессор П{number}"]
            self.assertEqual(source.sheet_state, "visible")
            self.assertEqual(source["A6"].value, pd.Timestamp(2026, number, 1))
            self.assertAlmostEqual(source["E6"].value, 84.01)
            self.assertAlmostEqual(source["F6"].value, 13.1)
            self.assertEqual(source["G6"].value, (number - 1) * 7 + 2)
            self.assertIn(f"'Компрессор П{number}'!$E$6:$E$7", chart.series[0].yVal.numRef.f)
            self.assertIn(f"'Компрессор П{number}'!$F$6:$F$7", chart._charts[1].series[0].yVal.numRef.f)
            labels = [item for item in chart.series if item.dLbls is not None]
            self.assertEqual(len(labels), 3)
            for label in labels:
                self.assertEqual(label.dLbls.spPr.solidFill.srgbClr, "FFFFFF")
                self.assertFalse(label.dLbls.showSerName)
                self.assertIsNotNone(label.xVal.numLit)
            self.assertIn("10,1 °С ≤ tнк ≤ 10,4 °С", labels[-1].tx.v)
            self.assertNotIn("ηoi", labels[-1].tx.v)
            self.assertIn("ηoi = 84,0 %", labels[0].tx.v)
            self.assertIsNotNone(chart.series[0].trendline)
            self.assertIsNotNone(chart._charts[1].series[0].trendline)
            self.assertEqual(chart.graphical_properties.solidFill.srgbClr, "FFFFFF")
            self.assertFalse(chart.x_axis.delete)
            for part in chart._charts:
                axis = part.y_axis
                self.assertFalse(axis.delete)
                self.assertGreater(axis.majorUnit, 0)
                span = axis.scaling.max - axis.scaling.min
                self.assertLessEqual(span / axis.majorUnit, 10)
                self.assertLessEqual(span / axis.minorUnit, 50)
                self.assertEqual(axis.txPr.p[0].pPr.defRPr.solidFill.srgbClr, "000000")
                self.assertEqual(axis.txPr.p[0].pPr.defRPr.sz, 1600)
                values = part.series[0].yVal.numRef.numCache.pt
                self.assertTrue(all(axis.scaling.min <= point.v <= axis.scaling.max
                                    for point in values))
            self.assertIsNone(chart._charts[1].y_axis.majorGridlines)
            self.assertIsNone(chart._charts[1].y_axis.minorGridlines)
            for part in chart._charts:
                self.assertFalse(part.series[0].trendline.dispEq)
                self.assertFalse(part.series[0].trendline.dispRSqr)
        descriptions = [row[0].value for row in sheet.iter_rows(min_col=1, max_col=1)
                        if isinstance(row[0].value, str)]
        self.assertTrue(any("10.100 ≤ tнк ≤ 10.400" in text
                            for text in descriptions))
        self.assertTrue(all(cell.value is None for row in sheet.iter_rows(min_col=67, max_col=69)
                            for cell in row))
        self.assertAlmostEqual(sheet._charts[2].series[0].yVal.numRef.numCache.pt[0].v, 84.01)

    def test_eta_filter_preserves_unrounded_boundary_values_and_source(self):
        dates = pd.date_range("2026-01-01", periods=4, freq="h")
        values = pd.DataFrame({
            "date": dates, "vna": [86.2] * 4, "t": [10.0] * 4,
            "p": [990.0] * 4, "pi_k": [14.0] * 4,
            "eta_compressor": [84.0096, 84.0204, 84.89, 82.86],
        })
        original = values.copy(deep=True)
        periods = pd.DataFrame([{"period": 1, "date_start": dates[0], "date_end": dates[-1]}])
        table = pd.DataFrame([{
            "vna_interval": "86.0-87.0", "t_1": 10, "t_2": 10, "p_1": 990, "p_2": 990,
            "date_1": dates[0], "date_2": dates[1], "eta_1": 84.010, "eta_2": 84.020,
        }])
        table.attrs["period_number"] = 1
        plots = prepare_compressor_plots(values, periods, [table])
        self.assertEqual(plots[0].points.eta.tolist(), [84.0096, 84.0204])
        pd.testing.assert_frame_equal(values, original)

    def test_value_axes_cover_flat_single_and_wide_series(self):
        for eta, pi_k in (([85.0], [14.0]),
                          ([85.0, 85.0], [14.0, 14.0]),
                          ([82.0, 89.0], [12.5, 21.0])):
            with self.subTest(eta=eta):
                points = pd.DataFrame({
                    "date": pd.date_range("2026-01-01", periods=len(eta)),
                    "eta": eta, "pi_k": pi_k,
                })
                top = pd.Series({
                    "date_1": points.date.iloc[0], "date_2": points.date.iloc[-1],
                    "eta_1": eta[0], "eta_2": eta[-1],
                    "pi_k_1": pi_k[0], "pi_k_2": pi_k[-1],
                })
                plot = CompressorPlot(1, points, top, 86, 87, 10, 11, 990, 991, 2)
                book = Workbook()
                sheet = book.active
                add_compressor_charts(sheet, [plot], 1)
                chart = sheet._charts[0]
                for axis, values in ((chart.y_axis, eta),
                                     (chart._charts[1].y_axis, pi_k)):
                    self.assertLessEqual(axis.scaling.min, min(values))
                    self.assertGreaterEqual(axis.scaling.max, max(values))
                    self.assertGreater(axis.scaling.max, axis.scaling.min)
                    self.assertLessEqual(
                        (axis.scaling.max - axis.scaling.min) / axis.majorUnit, 10,
                    )
                left, right = chart.y_axis.scaling, chart._charts[1].y_axis.scaling
                eta_bottom = (min(eta) - left.min) / (left.max - left.min)
                pi_top = (max(pi_k) - right.min) / (right.max - right.min)
                self.assertGreater(eta_bottom, pi_top)


if __name__ == "__main__":
    unittest.main()
