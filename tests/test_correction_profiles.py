import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scr"))
from correction_profiles import (FIELDS, KEYS, LEGACY_CELLS, CorrectionProfile, load_profile,
    save_profile, resolve_profile, require_profile, MARKER)
from correction_sheet import read_correction_sheet
from result_export import DataSheetValues, save_result_workbook


def profile(unit="11", value=1.):
    return CorrectionProfile.from_dict({"station_id":"syz_tec","unit":unit,
                                        "coefficients":{field.key:value for field in FIELDS}})


class CorrectionProfileTest(unittest.TestCase):
    def test_schema_has_63_unique_fields_and_complete_source_guide(self):
        self.assertEqual(len(FIELDS),63)
        self.assertEqual(len(KEYS),63)
        self.assertEqual(len(LEGACY_CELLS),63)

    def test_number_format_shape_identity_and_method(self):
        values=profile().to_dict()
        values["coefficients"]["pressure_power.p1"]="-1,234E-5"
        self.assertEqual(CorrectionProfile.from_dict(values).values["pressure_power.p1"],-1.234e-5)
        self.assertEqual(profile(value=0).values["pressure_power.p0"],0)
        for field, bad in (("coefficients",{}),("method","unknown")):
            with self.subTest(field=field),self.assertRaises(ValueError):
                CorrectionProfile.from_dict({**values,field:bad})
        for bad in ("nan","inf",""):
            with self.assertRaises(ValueError):
                CorrectionProfile.from_dict({**values,"coefficients":{**values["coefficients"],"pressure_power.p1":bad}})
        with self.assertRaises(ValueError):
            CorrectionProfile.from_dict(values,"syz_tec","10")

    def test_saved_profiles_are_isolated_atomic_and_versioned(self):
        with tempfile.TemporaryDirectory() as folder,patch("correction_profiles.profile_path",return_value=Path(folder)/"profiles.json"):
            a=save_profile(profile("10",2))
            b=save_profile(profile("11",3))
            self.assertEqual(load_profile("syz_tec","10").values,profile("10",2).values)
            self.assertEqual(save_profile(profile("11",4)).revision,2)
            before=(Path(folder)/"profiles.json").read_bytes()
            with patch("correction_profiles.os.replace",side_effect=OSError("locked")),self.assertRaises(OSError):
                save_profile(profile("11",5))
            self.assertEqual((Path(folder)/"profiles.json").read_bytes(),before)
            with patch("correction_profiles.pd.ExcelFile",side_effect=AssertionError("Workbook inspected")):
                self.assertEqual(resolve_profile("does_not_exist.xlsx","syz_tec","11").values,profile("11",4).values)

    def test_incomplete_or_corrupt_store_does_not_fall_back_to_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"profiles.json"
            path.write_text('{"schema_version":1,"profiles":{"syz_tec|11":{}}}',encoding="utf-8")
            with patch("correction_profiles.profile_path",return_value=path),self.assertRaises(ValueError):
                resolve_profile("book.xlsx","syz_tec","11")
            self.assertIn('"syz_tec|11":{}',path.read_text())

    def test_embedded_profile_works_without_any_old_coefficient_addresses(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"renamed.xlsx"
            data=DataSheetValues(pd.DataFrame([[None,"Дата"]]),0)
            save_result_workbook(data,[],path,correction_profile=profile())
            with patch("correction_profiles.load_profile",return_value=None):
                self.assertEqual(require_profile(path,"syz_tec","11").values,profile().values)
                with self.assertRaises(ValueError):
                    require_profile(path,"syz_tec","10")
            book=load_workbook(path,read_only=True)
            self.assertEqual(book["Поправки"]["M669"].value,None)
            book.close()

    def test_missing_profile_does_not_use_legacy_coordinates(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"book.xlsx"
            book=Workbook()
            book.active.title="Данные"
            book.create_sheet("Поправки")["M669"]=123
            book.save(path)
            with patch("correction_profiles.load_profile",return_value=None),self.assertRaisesRegex(ValueError,"Нет именованного"):
                require_profile(path,"syz_tec","11")

    def test_ready_profile_uses_old_exhaust_location_only_in_working_copy(self):
        from test_gas_profiles import profile as legacy_values
        from correction_profiles import named_legacy_values
        from data_assembly import AssembledData,calculate_data_columns
        from openpyxl.utils import get_column_letter
        raw=pd.DataFrame({"B":[pd.Timestamp("2026-01-01")],"C":[100100.],"F":[10.],"G":[50.],
            "L":[200.],"M":[220.],"P":[12.],"X":[1.31],"W":[365.],"AC":[.8],"AD":[.8],
            "AH":[20.],"AM":[5100.],"AO":[85.],"AS":[4.]})
        values=named_legacy_values(legacy_values())
        values["outlet_loss.a0"],values["outlet_loss.b0"]=1.,0.
        data=AssembledData(raw.copy(),{},0,0,{},[])
        calculate_data_columns(data,values)
        expected=data.values.at[0,"AQ"]
        letters=[get_column_letter(i) for i in range(1,51)]
        table=data.values.reindex(columns=letters)
        table["AC"]=float("nan")
        headers=["Дата" if letter=="B" else None if letter in ("A","AD") else letter for letter in letters]
        book_data=DataSheetValues(pd.DataFrame([headers,*table.itertuples(index=False,name=None)]),0)
        book_data.apply_correction_profile(values)
        self.assertAlmostEqual(book_data.values.iat[1,42],expected)
        self.assertTrue(pd.isna(book_data.values.iat[1,28]))
        self.assertEqual(book_data.values.iat[1,29],.8)


class CorrectionStyleTest(unittest.TestCase):
    def source(self,path):
        book=Workbook()
        book.active.title="Данные"
        for i in range(100):
            cell=book.active.cell(i+1,1,i)
            cell.font=Font(name="Arial",size=8+i/10)
            _=cell.style_id
        sheet=book.create_sheet("Поправки")
        sheet["M669"],sheet["N669"]=.01,.02
        sheet.column_dimensions["M"].font=Font(name="Courier New",size=17,bold=True)
        sheet.row_dimensions[851].number_format="0.000000"
        sheet.row_dimensions[851].fill=PatternFill("solid",fgColor="FFFF00")
        book.save(path)

    def test_dimension_styles_are_registered_in_output_not_copied_as_ids(self):
        with tempfile.TemporaryDirectory() as folder:
            source,target=Path(folder)/"source.xlsx",Path(folder)/"target.xlsx"
            self.source(source)
            snapshot=read_correction_sheet(source)
            save_result_workbook(DataSheetValues(pd.DataFrame([[None,"Дата"]]),0),[],target,correction_sheet=snapshot)
            book=load_workbook(target)
            sheet=book["Поправки"]
            self.assertEqual(sheet.column_dimensions["M"].font.name,"Courier New")
            self.assertEqual(sheet.column_dimensions["M"].font.sz,17)
            self.assertEqual(sheet.row_dimensions[851].number_format,"0.000000")
            self.assertEqual(sheet["M669"].value,.01)
            book.close()
            self.assertIsNotNone(read_correction_sheet(target))

    def test_previously_broken_dimension_references_are_recovered_in_memory(self):
        with tempfile.TemporaryDirectory() as folder:
            source=Path(folder)/"source.xlsx"
            self.source(source)
            with ZipFile(source) as z:
                parts={name:z.read(name) for name in z.namelist()}
            root=ET.fromstring(parts["xl/worksheets/sheet2.xml"])
            ns="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            root.find(f"{ns}cols/{ns}col").set("style","99999")
            parts["xl/worksheets/sheet2.xml"]=ET.tostring(root)
            with ZipFile(source,"w",ZIP_DEFLATED) as z:
                for name,value in parts.items():z.writestr(name,value)
            before=source.read_bytes()
            snapshot=read_correction_sheet(source)
            self.assertEqual(source.read_bytes(),before)
            book=load_workbook(io.BytesIO(snapshot.content))
            self.assertEqual(book["Поправки"]["M669"].value,.01)
            book.close()
