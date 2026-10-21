"""Проверка сводки деградации и выбора точек для диаграммы Excel."""

import io
import sys
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scr"))
from degradation_chart import prepare_degradation_plot  # noqa: E402
from pressure_chart import PressureChartData, write_pressure_sheet  # noqa: E402
from RH_raschet import build_weighted_delta_summary  # noqa: E402


class DegradationSummaryChartTest(unittest.TestCase):
    def test_weighted_summary_counts_only_power_loss_and_fuel_growth(self):
        table = pd.DataFrame({
            "Период": [1, 2, 3],
            "Часы начала": [0.0, 10.0, 30.0],
            "Часы конца": [10.0, 30.0, 60.0],
            "ΔNприв": [2.0, -4.0, -2.0],
            "Δbприв": [-2.0, 4.0, 1.0],
            "Δη, п.п.": [0.0, -0.1, -0.2],
        })
        table.attrs["period_start_values"] = {1: {"N": 100.0, "b": 200.0}}
        original = table.copy(deep=True)

        summary = build_weighted_delta_summary(table)

        self.assertAlmostEqual(summary["delta_n"], 140.0 / 60.0)
        self.assertAlmostEqual(summary["delta_b"], 110.0 / 60.0)
        self.assertAlmostEqual(summary["delta_n_percent"], 140.0 / 60.0)
        self.assertAlmostEqual(summary["delta_b_percent"], 110.0 / 60.0 / 2.0)
        pd.testing.assert_frame_equal(table, original)

    def test_chart_uses_each_periods_own_vna_interval(self):
        records = []
        for period, base, interval in (
            (1, "2026-01-01", 86),
            (2, "2026-02-01", 85),
            (3, "2026-03-01", 87),
        ):
            for index, offset in enumerate((0.2, 0.4, -0.7)):
                records.append({
                    "period": period,
                    "date": pd.Timestamp(base) + pd.Timedelta(minutes=30 * index),
                    "vna": interval + offset,
                    "power": 100 - period - index,
                    "fuel_specific": 280 + period + index,
                    "eta_gtu": 28 - period,
                    "vlo": 100 + index,
                    "fgo_fto": 200 + index,
                    "power_n": 80 + index,
                })
        values = pd.DataFrame(records)
        periods = pd.DataFrame([
            {
                "period": number,
                "date_start": group["date"].min(),
                "date_end": group["date"].max(),
            }
            for number, group in values.groupby("period")
        ])
        deltas = pd.DataFrame({
            "Период": [1, 2, 3],
            "ΔNприв": [2.0, -4.0, -3.0],
            "Δbприв": [-1.0, 2.0, 3.0],
            "Δη, п.п.": [0.0, -0.1, -0.2],
        })
        deltas.attrs["vna_intervals"] = {
            1: (86.0, 87.0, 2),
            2: (85.0, 86.0, 2),
            3: (87.0, 88.0, 2),
        }

        plot = prepare_degradation_plot(values, periods, deltas)
        self.assertEqual(len(plot.points), 6)
        self.assertEqual(
            [(p["period"], p["vna_start"], p["count"]) for p in plot.periods],
            [(1, 86.0, 2), (2, 85.0, 2), (3, 87.0, 2)],
        )

        book = Workbook(write_only=True)
        source_sheet = book.create_sheet("Данные")
        headers = ["Дата", "ВЛО, Па", "ФГО+ФТО, Па", "N, МВт"]
        source_sheet.append(headers)
        for row in values.itertuples():
            source_sheet.append([
                row.date.to_pydatetime(), row.vlo, row.fgo_fto, row.power_n,
            ])
        write_pressure_sheet(
            book, headers, 0, len(values) + 1,
            PressureChartData(values, periods, None, (), values, deltas),
        )
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        saved = load_workbook(output)
        chart_sheet = saved["Графики"]
        self.assertEqual(len(chart_sheet._charts), 2)
        degradation_chart = chart_sheet._charts[1]
        self.assertEqual(len(degradation_chart._charts), 2)
        self.assertEqual(len(degradation_chart.series), 3)
        self.assertEqual(len(degradation_chart._charts[1].series), 1)
        self.assertEqual(chart_sheet["AM2"].value, values.iloc[0]["date"])
        self.assertEqual(chart_sheet["AQ2"].value, 86.2)

        # Отсутствие каналов перепада не должно убирать диаграмму деградации.
        without_pressure = Workbook(write_only=True)
        without_pressure.create_sheet("Данные").append(["Дата"])
        write_pressure_sheet(
            without_pressure, ["Дата"], 0, len(values) + 1,
            PressureChartData(values, periods, None, (), values, deltas),
        )
        standalone_output = io.BytesIO()
        without_pressure.save(standalone_output)
        standalone_output.seek(0)
        standalone = load_workbook(standalone_output)
        self.assertEqual(len(standalone["Графики"]._charts), 1)


if __name__ == "__main__":
    unittest.main()
