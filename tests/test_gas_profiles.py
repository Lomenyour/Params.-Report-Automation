"""Паспорт, перенос поправок и повторный расчёт без исходной книги/кеша."""
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scr"))
from gas_passport import (GasPassport, MARKER, load_passport, save_passport,
                          read_book_passport, resolve_passport)
from gas_passport_ui import passport_from_fields
from correction_sheet import read_correction_sheet
from data_assembly import (read_calculation_profile, recalculate_gas_columns,
                           AssembledData, calculate_data_columns)
from result_export import DataSheetValues, save_result_workbook
from correction_profiles import CorrectionProfile, KEYS, named_legacy_values, LEGACY_CELLS


def passport(unit="11", density=.7, q=34.13):
    return GasPassport.from_dict({"station_id": "syz_tec", "unit": unit,
        "density_kg_per_nm3": density, "q_lower_mj_per_nm3": q,
        "carbon_hydrogen_ratio": 3.066677,
        "components": {"Methane": .9631, "Nitrogen": .0369}})


def profile():
    result = dict.fromkeys(LEGACY_CELLS, 0.)
    result.update({"N669": 1., "AF788": 1., "AD888": 1., "P712": 1.,
        "T507": 1., "AE969": 1., "U1107": 1., "T1019": 1., "AX851": 1.,
        "M605": .05, "O605": 100., "gas_c_h": 3., "gas_density": .7,
        "qnr_denominator": .7, "gas_q_lower": 34.13})
    return result


def make_book(path):
    book = Workbook()
    data = book.active
    data.title = "Данные"
    data.append([None, "Дата"] + [f"Параметр {i}" for i in range(2, 50)])
    data.append([0, pd.Timestamp("2026-01-01").to_pydatetime()] + [1.] * 48)
    corrections = book.create_sheet("Поправки")
    for key, value in profile().items():
        if key in LEGACY_CELLS:
            corrections[key] = value
    corrections["A1"] = "Исходные поправки"
    corrections["A1"].font = Font(name="Times New Roman", size=16, bold=True)
    corrections["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    corrections.merge_cells("A1:C1")
    corrections.column_dimensions["A"].width = 35
    corrections.row_dimensions[1].height = 32
    corrections["A2"], corrections["B2"], corrections["A3"], corrections["B3"] = 1, 2, 3, 4
    corrections["A4"] = "=SUM(A2:A3)"
    chart = ScatterChart()
    chart.series.append(Series(Reference(corrections, min_col=2, min_row=2, max_row=3),
                               Reference(corrections, min_col=1, min_row=2, max_row=3)))
    corrections.add_chart(chart, "D2")
    gas = book.create_sheet("Результаты")
    # Намеренно не AC:AE, чтобы проверить распознавание газовых заголовков.
    gas.append(["C/H", "ρ, кг/м3", "Qнр, МДж/м3", .9631, .0369])
    gas.append([3.066677, .7, 34.13, "Methane", "Nitrogen"])
    book.save(path)
    # Известный сохранённый результат формулы, как в настоящем Excel.
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    members["xl/worksheets/sheet2.xml"] = members["xl/worksheets/sheet2.xml"].replace(
        b'<f>SUM(A2:A3)</f><v></v>', b'<f>SUM(A2:A3)</f><v>4</v>')
    with ZipFile(path, "w", ZIP_DEFLATED) as output:
        for name, content in members.items():
            output.writestr(name, content)


def corrections():
    values = named_legacy_values(profile())
    return CorrectionProfile.from_dict({"station_id":"syz_tec","unit":"11",
                                        "coefficients":{key:values[key] for key in KEYS}})


class GasPassportTest(unittest.TestCase):
    def test_decimal_input_and_composition_are_not_normalized(self):
        value = passport_from_fields({"density_kg_per_nm3": "0,7", "q_lower_mj_per_nm3": "34,13",
            "carbon_hydrogen_ratio": "3,066677", "component:Methane": "0,9631",
            "component:Nitrogen": "0,00716", "component:Helium": ""}, "syz_tec", "11")
        self.assertEqual(dict(value.components)["Methane"], .9631)
        self.assertTrue(value.composition_warning)
        self.assertAlmostEqual(value.q_lower_mj_per_kg, 34.13 / .7)

    def test_bad_values_and_wrong_unit_are_rejected(self):
        for field, bad in (("density_kg_per_nm3", 0), ("q_lower_mj_per_nm3", "nan"),
                           ("components", {"Methane": 96.31})):
            with self.subTest(field=field), self.assertRaises(ValueError):
                GasPassport.from_dict({**passport().to_dict(), field: bad})
        with self.assertRaises(ValueError):
            GasPassport.from_dict(passport().to_dict(), "syz_tec", "10")

    def test_store_is_separate_per_unit_and_atomic_on_failed_replace(self):
        with tempfile.TemporaryDirectory() as folder, patch("gas_passport.passport_path", return_value=Path(folder) / "gas.json"):
            save_passport(passport("10", .71))
            save_passport(passport("11", .7))
            self.assertEqual(load_passport("syz_tec", "10").density_kg_per_nm3, .71)
            original = (Path(folder) / "gas.json").read_bytes()
            with patch("gas_passport.os.replace", side_effect=OSError("locked")), self.assertRaises(OSError):
                save_passport(passport("11", .72))
            self.assertEqual((Path(folder) / "gas.json").read_bytes(), original)
            self.assertIsNone(load_passport("other", "11"))

    def test_corrupt_store_is_not_silently_replaced(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "gas.json"
            path.write_text("bad", encoding="utf-8")
            with patch("gas_passport.passport_path", return_value=path), self.assertRaises(ValueError):
                save_passport(passport())
            self.assertEqual(path.read_text(), "bad")

    def test_user_passport_has_priority_over_older_book(self):
        with tempfile.TemporaryDirectory() as folder, patch("gas_passport.passport_path", return_value=Path(folder) / "gas.json"):
            path = Path(folder) / "СызТЭЦ ГТ-11.xlsx"
            make_book(path)
            self.assertEqual(read_book_passport(path, "syz_tec", "11").q_lower_mj_per_nm3, 34.13)
            save_passport(passport(q=35.))
            self.assertEqual(resolve_passport(path, "syz_tec", "11").q_lower_mj_per_nm3, 35.)

    def test_passport_block_is_selected_instead_of_repeated_calculation_header(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "СызТЭЦ ГТ-11.xlsx"
            make_book(path)
            book = load_workbook(path)
            sheet = book["Результаты"]
            sheet["K1"], sheet["K2"] = "Qнр, МДж/м3", 1000
            book.save(path)
            book.close()
            self.assertEqual(read_book_passport(path, "syz_tec", "11").q_lower_mj_per_nm3, 34.13)


class CorrectionRoundTripTest(unittest.TestCase):
    def test_export_preserves_entire_correction_sheet_and_loads_without_cache(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "СызТЭЦ ГТ-11.xlsx"
            target = Path(folder) / "СызТЭЦ ГТ-11_результаты.xlsx"
            make_book(source)
            before = source.read_bytes()
            snapshot = read_correction_sheet(source)
            values = DataSheetValues(pd.DataFrame([[None, "Дата"], [0, pd.Timestamp("2026-01-01")]]), 0)
            save_result_workbook(values, [], target, correction_sheet=snapshot, gas_passport=passport(), correction_profile=corrections())
            book = load_workbook(target)
            sheet = book["Поправки"]
            self.assertEqual(sheet["A4"].value, 4)
            self.assertEqual(sheet["A1"].font.name, "Times New Roman")
            self.assertEqual(sheet["A1"].font.sz, 16)
            self.assertEqual(sheet.column_dimensions["A"].width, 35)
            self.assertEqual(sheet.row_dimensions[1].height, 32)
            self.assertIn("A1:C1", {str(r) for r in sheet.merged_cells.ranges})
            self.assertEqual(len(sheet._charts), 1)
            self.assertNotIn("Результаты", book.sheetnames)
            self.assertEqual(book.sheetnames[0], "Данные")
            book.close()
            source.unlink()  # Для повторного чтения доступна только сформированная книга.
            with patch("correction_profiles.load_profile", return_value=None):
                actual = read_calculation_profile(target)
            for key in KEYS:
                self.assertEqual(actual[key], corrections().values[key])
            self.assertEqual(actual["gas_q_lower"], 34.13)
            self.assertEqual(read_book_passport(target, "syz_tec", "11"), passport())

    def test_snapshot_is_immutable_and_second_export_has_one_passport(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "СызТЭЦ ГТ-11.xlsx"
            first = Path(folder) / "СызТЭЦ ГТ-11_1.xlsx"
            second = Path(folder) / "СызТЭЦ ГТ-11_2.xlsx"
            make_book(source)
            snapshot = read_correction_sheet(source)
            changed = load_workbook(source)
            changed["Поправки"]["M605"] = 10000
            changed.save(source)
            changed.close()
            values = DataSheetValues(pd.DataFrame([[None, "Дата"]]), 0)
            save_result_workbook(values, [], first, correction_sheet=snapshot, gas_passport=passport())
            save_result_workbook(values, [], second, correction_sheet=read_correction_sheet(first), gas_passport=passport(q=35.))
            book = load_workbook(second, data_only=True)
            sheet = book["Поправки"]
            self.assertEqual(sheet["M605"].value, .05)
            self.assertEqual(sum(row[0].value == MARKER for row in sheet), 1)
            book.close()
            self.assertEqual(read_book_passport(second, "syz_tec", "11").q_lower_mj_per_nm3, 35.)

    def test_ready_and_import_gas_calculation_agree_and_preserve_raw_columns(self):
        raw = pd.DataFrame({"C": [100100., 99000.], "F": [10., 17.], "G": [50., 40.],
            "L": [200., 300.], "M": [220., 320.], "P": [12., 19.], "X": [1.31, 1.42],
            "W": [365., 388.], "AC": [.8, 1.1], "AH": [20., 30.], "AM": [5100., 5150.],
            "AO": [85., 90.], "AS": [4., 5.]})
        data = AssembledData(raw.copy(), {}, 0, 0, {}, [])
        calculated = calculate_data_columns(data, named_legacy_values(profile())).values
        same = calculated.copy(deep=True)
        recalculate_gas_columns(same, named_legacy_values(profile()))
        pd.testing.assert_frame_equal(calculated, same)
        pd.testing.assert_frame_equal(raw, calculated[raw.columns])
        self.assertAlmostEqual(calculated.loc[0, "AK"], 34.13 / .7)
        self.assertAlmostEqual(calculated.loc[0, "AT"], 4 * 3600 / .7)
        changed = calculated.copy(deep=True)
        recalculate_gas_columns(changed, {**named_legacy_values(profile()), "gas_density": .8, "qnr_denominator": .8})
        self.assertNotEqual(calculated.loc[0, "AV"], changed.loc[0, "AV"])

    def test_extra_notes_after_exported_passport_are_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "СызТЭЦ ГТ-11.xlsx"
            first, second = Path(folder) / "first.xlsx", Path(folder) / "second.xlsx"
            make_book(source)
            data = DataSheetValues(pd.DataFrame([[None, "Дата"]]), 0)
            save_result_workbook(data, [], first, correction_sheet=read_correction_sheet(source), gas_passport=passport())
            book = load_workbook(first)
            book["Поправки"]["A1500"] = "Дополнительное примечание пользователя"
            book.save(first)
            book.close()
            save_result_workbook(data, [], second, correction_sheet=read_correction_sheet(first), gas_passport=passport())
            check = load_workbook(second, read_only=True)
            self.assertEqual(check["Поправки"]["A1500"].value, "Дополнительное примечание пользователя")
            check.close()


if __name__ == "__main__":
    unittest.main()
