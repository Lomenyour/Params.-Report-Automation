"""Выбор из существующих трёх вариантов и нативная РХ с верными датами."""
import copy
import io
from pathlib import Path
import sys
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scr"))
from RH_raschet import select_rh_variants, plot_selected_windows
from rh_chart import prepare_rh_plots, add_rh_charts


class RHSingleChartTest(unittest.TestCase):
    def candidate(self, rank, end_slope, period=1):
        first = pd.Timestamp(2026, period, 1)
        return {"period": period, "rank": rank, "k_start": 200., "k_end": end_slope,
                "b_start": 6500., "b_end": 6300.,
                "start_window_begin": first, "start_window_end": first + pd.Timedelta(hours=23.5),
                "end_window_begin": first + pd.Timedelta(days=2),
                "end_window_end": first + pd.Timedelta(days=2, hours=23.5)}

    def values(self):
        result = []
        for start, k, b in (("2026-01-01", 200., 6500.), ("2026-01-03", 201., 6300.)):
            for index, date in enumerate(pd.date_range(start, periods=48, freq="30min")):
                power = 65. + index * .4
                result.append({"date": date, "period": 1, "power": power,
                               "fuel": k * power + b, "excel_row": len(result) + 2})
        return pd.DataFrame(result)

    def test_selection_first_second_third_equal_fallback_and_short_lists(self):
        for slopes, rank in (([201., 202., 203.], 1), ([199., 200., 201.], 2),
                             ([199., 198., 201.], 3), ([199., 198., 197.], 1),
                             ([200.], 1), ([199., 201.], 2)):
            with self.subTest(slopes=slopes):
                rows = [self.candidate(i + 1, slope) for i, slope in enumerate(slopes)]
                original = copy.deepcopy(rows)
                self.assertEqual(select_rh_variants(rows)[0]["rank"], rank)
                self.assertEqual(rows, original)
        self.assertEqual(select_rh_variants([]), [])

    def test_each_period_selected_independently_without_line_position_criterion(self):
        rows = [self.candidate(1, 199.), self.candidate(2, 201.),
                self.candidate(1, 199., 2), self.candidate(2, 198., 2)]
        selected = select_rh_variants(rows)
        self.assertEqual([(row["period"], row["rank"]) for row in selected], [(1, 2), (2, 1)])
        # END с большим k может быть ниже из-за меньшего b: это не отменяет выбор.
        self.assertLess(selected[0]["k_end"] * 70 + selected[0]["b_end"], 200 * 70 + 6500)

    def test_excel_and_gui_use_same_candidate_and_preserve_window_values(self):
        rows = [self.candidate(1, 199.), self.candidate(2, 201.), self.candidate(3, 202.)]
        values = self.values()
        original = values.copy(deep=True)
        plots = prepare_rh_plots(values, rows)
        self.assertEqual(len(plots), 1)
        self.assertEqual(plots[0].result["rank"], 2)
        figure = plot_selected_windows(values, rows, show=False)
        self.addCleanup(plt.close, figure)
        self.assertEqual(len([axis for axis in figure.axes if axis.has_data()]), 1)
        self.assertIn("вариант №2", figure.axes[0].get_title())
        self.assertEqual(figure.axes[0].xaxis.label.get_size(), 16)
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Графики")
        cells, bottom = add_rh_charts(sheet, plots, 1)
        for i in range(1, bottom + 1):
            sheet.append([cells[i][1][0]] if i in cells else [])
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        with ZipFile(output) as archive:
            ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
                  "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            chart = ET.fromstring(archive.read("xl/charts/chart1.xml"))
            # DrawingML разрешает только один способ заливки линии.
            # Excel отвергает сочетание noFill и solidFill, хотя XML читается.
            for line in chart.findall(".//a:ln", ns):
                fills = [line.find(f"a:{kind}", ns) for kind in
                         ("noFill", "solidFill", "gradFill", "pattFill")]
                self.assertLessEqual(sum(fill is not None for fill in fills), 1)
            labels = ["".join(label.itertext()) for label in chart.findall(".//c:dLbl/c:tx/c:rich", ns)]
            self.assertIn("Условно-чистое состояние 01.01.2026", labels[1])
            self.assertIn("Условно-загрязнённое состояние 03.01.2026", labels[0])
            for label in chart.findall(".//c:dLbl", ns):
                self.assertIsNone(label.find("c:layout", ns))
                self.assertEqual(label.find("c:tx/c:rich/a:p/a:pPr/a:defRPr", ns).get("sz"), "1000")
        output.seek(0)
        saved = load_workbook(output)
        data = saved["Данные РХ"]
        self.assertEqual(data["C3"].value, values.iloc[0].power)
        self.assertEqual(data["D3"].value, values.iloc[0].fuel)
        self.assertEqual(data["E3"].value, values.iloc[0].excel_row)
        self.assertEqual(len(saved["Графики"]._charts), 1)
        chart = saved["Графики"]._charts[0]
        self.assertEqual(chart.anchor.to.col - chart.anchor._from.col, 14)
        self.assertEqual(chart.anchor.to.row - chart.anchor._from.row, 33)
        self.assertEqual(chart.anchor.editAs, "twoCell")
        pd.testing.assert_frame_equal(values, original)

    def test_three_periods_share_sheet_without_crossing_data_or_curve_ranges(self):
        parts, candidates = [], []
        for number in (1, 2, 3):
            part = self.values()
            part["period"] = number
            part["date"] += pd.DateOffset(months=number - 1)
            part["fuel"] += number * 1000
            parts.append(part)
            candidates.append(self.candidate(1, 201., number))
        values = pd.concat(parts, ignore_index=True)
        original = values.copy(deep=True)
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Графики")
        cells, bottom = add_rh_charts(sheet, prepare_rh_plots(values, candidates), 1)
        for _ in range(bottom):
            sheet.append([])
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        saved = load_workbook(output)
        self.assertEqual(saved.sheetnames, ["Графики", "Данные РХ"])
        source = saved["Данные РХ"]
        for number, chart in enumerate(saved["Графики"]._charts, 1):
            first = 3 + (number - 1) * 104
            self.assertIn(f"'Данные РХ'!$D${first}:$D${first+47}", chart.series[0].yVal.numRef.f)
            self.assertIn(f"'Данные РХ'!$D${first+48}:$D${first+95}", chart.series[1].yVal.numRef.f)
            self.assertIn(f"'Данные РХ'!$H${first}:$H${first+99}", chart.series[2].yVal.numRef.f)
            self.assertEqual(source[f"B{first}"].value, pd.Timestamp(2026, number, 1))
            self.assertEqual(source[f"D{first}"].value, parts[number-1].iloc[0].fuel)
            self.assertAlmostEqual(source[f"H{first}"].value,
                candidates[number-1]["k_start"] * source[f"G{first}"].value + candidates[number-1]["b_start"])
        pd.testing.assert_frame_equal(values, original)


if __name__ == "__main__":
    unittest.main()
